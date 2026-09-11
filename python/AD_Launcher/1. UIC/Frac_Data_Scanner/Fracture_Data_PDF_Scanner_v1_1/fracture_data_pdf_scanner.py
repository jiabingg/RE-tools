#!/usr/bin/env python3
"""
Fracture Data PDF Scanner
=========================

A single-file Tkinter GUI that scans folders of PDF well files and flags
files containing likely:

- Step-rate tests
- Mini-frac / minifrac tests
- DFITs (diagnostic fracture injection tests)
- Leak-off / formation-integrity tests
- Hydraulic-fracture treatment records
- Fracture-gradient, ISIP, breakdown-pressure, closure-pressure,
  treating-pressure, rate, and proppant data

The scanner uses embedded PDF text first. For image-only scanned pages it can
optionally run OCR with Tesseract. Results can be exported to Excel or CSV.

The included bootstrap launcher creates a local .venv, verifies the required
packages, installs missing dependencies, checks for Tesseract OCR, and starts
the application with the managed virtual environment.

Use Start_Fracture_Data_Scanner.bat on Windows.

Tested with Python 3.10+.
"""

from __future__ import annotations

import argparse
import csv
import concurrent.futures
import os
import queue
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import traceback
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable, Optional, Sequence

# GUI imports are part of the Python standard library on normal Windows builds.
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter import font as tkfont

try:
    import fitz  # PyMuPDF
except ImportError:  # handled cleanly when the app starts
    fitz = None  # type: ignore[assignment]

try:
    import pytesseract
except ImportError:
    pytesseract = None  # type: ignore[assignment]

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None  # type: ignore[assignment]
    Alignment = Font = PatternFill = get_column_letter = None  # type: ignore[assignment]


APP_TITLE = "Fracture Data PDF Scanner"
APP_VERSION = "1.1"
DEFAULT_OCR_DPI = 150
MIN_EMBEDDED_TEXT_ALNUM = 80
MAX_MATCHES_PER_RULE_PER_PAGE = 3
MAX_DETAILS_PER_FILE = 100
MAX_EXCEL_CELL = 32000

STATUS_RANK = {
    "YES - Test/treatment data": 5,
    "POSSIBLE - Review": 4,
    "POSSIBLE - Planned only": 3,
    "REFERENCE/ASSUMED ONLY": 2,
    "NO CLEAR DATA": 1,
    "ERROR": 0,
}


@dataclass(frozen=True)
class KeywordRule:
    category: str
    label: str
    pattern: re.Pattern[str]
    weight: int
    strong: bool = False


@dataclass
class MatchDetail:
    page: int
    source: str
    category: str
    matched_text: str
    snippet: str


@dataclass
class PageEvidence:
    page: int
    source: str
    categories: set[str] = field(default_factory=set)
    category_weights: dict[str, int] = field(default_factory=dict)
    details: list[MatchDetail] = field(default_factory=list)
    gradients: list[str] = field(default_factory=list)
    isip_values: list[str] = field(default_factory=list)
    closure_values: list[str] = field(default_factory=list)
    breakdown_values: list[str] = field(default_factory=list)
    treating_values: list[str] = field(default_factory=list)
    rate_values: list[str] = field(default_factory=list)
    proppant_values: list[str] = field(default_factory=list)
    actual_marker: bool = False
    planned_marker: bool = False
    assumed_marker: bool = False
    text_quality: int = 0


@dataclass
class ScanOptions:
    use_ocr: bool = True
    ocr_dpi: int = DEFAULT_OCR_DPI
    tesseract_path: str = ""
    stop_after_strong_evidence: bool = True


@dataclass
class ScanResult:
    file_path: str
    file_name: str
    status: str
    score: int
    api_numbers: str
    well_name: str
    categories: str
    fracture_gradients: str
    isip_values: str
    closure_pressures: str
    breakdown_pressures: str
    treating_pressures: str
    rate_values: str
    proppant_values: str
    matched_pages: str
    pages_scanned: int
    total_pages: int
    ocr_pages: int
    image_only_pages_not_ocrd: int
    notes: str
    details: list[MatchDetail] = field(default_factory=list)
    elapsed_seconds: float = 0.0

    def summary_row(self) -> dict[str, object]:
        return {
            "Status": self.status,
            "Score": self.score,
            "API Number(s)": self.api_numbers,
            "Well Name": self.well_name,
            "Categories": self.categories,
            "Fracture Gradient": self.fracture_gradients,
            "ISIP": self.isip_values,
            "Closure Pressure": self.closure_pressures,
            "Breakdown Pressure": self.breakdown_pressures,
            "Treating Pressure": self.treating_pressures,
            "Rate Data": self.rate_values,
            "Proppant Data": self.proppant_values,
            "Matched Pages": self.matched_pages,
            "Pages Scanned": self.pages_scanned,
            "Total Pages": self.total_pages,
            "OCR Pages": self.ocr_pages,
            "Image-only Pages Not OCR'd": self.image_only_pages_not_ocrd,
            "Notes": self.notes,
            "Elapsed Seconds": round(self.elapsed_seconds, 1),
            "File Name": self.file_name,
            "File Path": self.file_path,
        }


# ---------------------------- Keyword definitions ---------------------------


def _rx(pattern: str) -> re.Pattern[str]:
    return re.compile(pattern, re.IGNORECASE | re.MULTILINE)


KEYWORD_RULES: tuple[KeywordRule, ...] = (
    KeywordRule(
        "Step-rate test",
        "step-rate test",
        _rx(r"\bstep[\s\-_/]*rate(?:\s+(?:injection\s+)?tests?)?\b"),
        9,
        True,
    ),
    KeywordRule(
        "Mini-frac / DFIT",
        "mini-frac",
        _rx(r"\b(?:mini[\s\-_/]*frac(?:ture)?|data[\s\-_/]*frac|fracture\s+calibration)(?:\s+(?:tests?|jobs?|treatments?))?\b"),
        9,
        True,
    ),
    KeywordRule(
        "Mini-frac / DFIT",
        "DFIT",
        _rx(r"\b(?:d\s*f\s*i\s*t|diagnostic\s+fracture\s+injection\s+tests?)\b"),
        9,
        True,
    ),
    KeywordRule(
        "Leak-off / FIT",
        "leak-off or formation-integrity test",
        _rx(
            r"\b(?:leak[\s\-]*off\s+(?:tests?|pressure)|formation\s+integrity\s+tests?|"
            r"lot\s+test|fit\s+test)\b"
        ),
        7,
        True,
    ),
    KeywordRule(
        "Fracture gradient",
        "fracture gradient",
        _rx(
            r"\b(?:formation\s+)?frac(?:ture)?\s+grad(?:ient)?\b|"
            r"\bfracture[\s\-]*gradient\b"
        ),
        8,
        True,
    ),
    KeywordRule(
        "Breakdown pressure",
        "breakdown pressure",
        _rx(
            r"\b(?:formation\s+)?break[\s\-]*down(?:\s+pressure)?\b|"
            r"\bformation\s+breakdown\b"
        ),
        6,
        True,
    ),
    KeywordRule(
        "Closure pressure",
        "closure pressure",
        _rx(r"\b(?:fracture\s+)?closure\s+(?:pressure|stress)\b"),
        6,
        True,
    ),
    KeywordRule(
        "ISIP",
        "ISIP",
        _rx(
            r"\b(?:i\s*s\s*i\s*p|i\s*s\s*p|instantaneous\s+shut[\s\-]*in\s+pressure)\b"
        ),
        6,
        True,
    ),
    KeywordRule(
        "Hydraulic-fracture treatment",
        "hydraulic-fracture treatment",
        _rx(
            r"\b(?:hydraulic(?:ally)?\s+fractur(?:e|ed|ing)|"
            r"fracture\s+stimulat(?:e|ed|ion)|frac(?:ture)?\s+(?:treatment|job|stimulation)|"
            r"stimulat(?:e|ed|ion)\s+(?:of\s+)?(?:the\s+)?[a-z0-9\-]+\s+formation)\b"
        ),
        6,
        True,
    ),
    KeywordRule(
        "Treating pressure",
        "treating pressure",
        _rx(
            r"\b(?:treating|treatment)\s+press(?:ure)?\b|"
            r"\b(?:max(?:imum)?\s*/\s*min(?:imum)?\s*/\s*avg|max\s*/\s*min\s*/\s*average)\s+tp\b|"
            r"\bmax\s*/\s*min\s*/\s*avg\s+tp\b"
        ),
        4,
        False,
    ),
    KeywordRule(
        "Pressure-rate diagnostics",
        "pressure-rate plot or diagnostics",
        _rx(
            r"\b(?:pressure\s*(?:vs\.?|versus|/|-to-)\s*rate|"
            r"rate\s*(?:vs\.?|versus|/|-to-)\s*pressure|"
            r"resultant\s+plot|fracture\s+point|pressure\s+buildup\s+profile|"
            r"diagnostic(?:s)?)\b"
        ),
        3,
        False,
    ),
    KeywordRule(
        "Rate data",
        "injection or pump rate",
        _rx(r"\b(?:injection|surface\s+injection|pump(?:ing)?)\s+rate\b"),
        2,
        False,
    ),
    KeywordRule(
        "Proppant data",
        "proppant or sand concentration",
        _rx(r"\b(?:proppant|mesh\s+sand|ppg\s+(?:sand|sd)|sand\s+concentration)\b"),
        2,
        False,
    ),
)

# Plausible fracture gradients usually fall below 2 psi/ft. OCR often drops a
# leading decimal, so the extractor keeps the raw value and adds a warning.
GRADIENT_RE = _rx(
    r"(?<![\d.])(?P<value>\d{1,4}(?:[.,]\d{1,4})?|[.,]\d{1,4})\s*"
    r"(?P<unit>psi\s*(?:/|per)\s*(?:ft|foot|feet)|psig\s*(?:/|per)\s*(?:ft|foot|feet))"
)

NUMBER_VALUE = r"(?:\d{1,3}(?:,\d{3})+|\d{2,6})(?:\.\d+)?"
PRESSURE_NUMBER = rf"(?P<value>{NUMBER_VALUE})\s*(?:psi|psig)\b"

ISIP_PATTERNS = (
    _rx(
        r"\b(?:i\s*s\s*i\s*p|i\s*s\s*p|instantaneous\s+shut[\s\-]*in\s+pressure)\b"
        r"[^\d]{0,30}" + PRESSURE_NUMBER
    ),
    _rx(
        PRESSURE_NUMBER
        + r"[^a-z0-9]{0,20}\b(?:i\s*s\s*i\s*p|i\s*s\s*p|instantaneous\s+shut[\s\-]*in\s+pressure)\b"
    ),
)

CLOSURE_PATTERNS = (
    _rx(r"\b(?:fracture\s+)?closure\s+(?:pressure|stress)\b[^\d]{0,40}" + PRESSURE_NUMBER),
    _rx(PRESSURE_NUMBER + r"[^a-z0-9]{0,25}\b(?:fracture\s+)?closure\s+(?:pressure|stress)\b"),
)

BREAKDOWN_PATTERNS = (
    _rx(
        r"\b(?:formation\s+)?break[\s\-]*down(?:\s+pressure)?\b[^\d]{0,40}"
        + PRESSURE_NUMBER
    ),
    _rx(
        PRESSURE_NUMBER
        + r"[^a-z0-9]{0,25}\b(?:formation\s+)?break[\s\-]*down(?:\s+pressure)?\b"
    ),
)

TREATING_PATTERNS = (
    _rx(
        r"\b(?:max(?:imum)?\s*/\s*min(?:imum)?\s*/\s*(?:avg|average)|"
        r"max\s*/\s*min\s*/\s*avg)\s+(?:treating\s+pressure|tp)\s*[:=]?\s*"
        r"(?P<values>[\d,./\s]+)\s*(?:psi|psig)\b"
    ),
    _rx(r"\b(?:treating|treatment)\s+pressure\b[^\d]{0,40}" + PRESSURE_NUMBER),
    _rx(PRESSURE_NUMBER + r"[^a-z0-9]{0,25}\b(?:treating|treatment)\s+pressure\b"),
)

RATE_RE = _rx(
    r"(?<!\w)(?P<value>\d{1,6}(?:\.\d+)?)\s*"
    r"(?P<unit>bpm|gpm|bpd|bwpd|bspd|bbls?\s*/\s*min|bbl\s*/\s*min)\b"
)

PROPPANT_RE = _rx(
    r"(?P<amount>\d{1,3}(?:,\d{3})+|\d{4,7})\s*(?:lbs?|pounds?|[l1it][b8])\b"
    r".{0,45}?\b(?P<type>\d{1,3}\s*/\s*\d{1,3}\s+mesh\s+sand|sand|proppant)\b"
)

ACTUAL_MARKER_RE = _rx(
    r"\b(?:conducted|performed|test\s+was\s+performed|pumped|pump\s+and\s+flush|"
    r"terminated\s+(?:the\s+)?treatment|actual\s+(?:test|treatment)|"
    r"maximum\s+treating\s+pressure|minimum\s+treating\s+pressure|"
    r"max\s*/\s*min\s*/\s*avg\s+tp|recorded|witnessed|flow\s*back|"
    r"frac\s+stimulated|fracture\s+stimulated)\b"
)

PLANNED_MARKER_RE = _rx(
    r"\b(?:proposed|planned|planning|program|procedure|recommendation|recommended|"
    r"will\s+be|to\s+be\s+(?:conducted|performed|pumped|fractured|stimulated)|"
    r"shall\s+be|anticipated|design\s+rate|design\s+pressure)\b"
)

ASSUMED_MARKER_RE = _rx(
    r"\b(?:assum(?:e|ed|ing|ption)|typical|default|estimated|conservative|"
    r"maximum\s+allowable|permitted|permit\s+limit|project\s+limit|"
    r"not\s+to\s+exceed|used\s+for\s+calculation|reference\s+value)\b"
)


# ------------------------------- Text helpers -------------------------------


def normalize_text(text: str) -> str:
    """Normalize common OCR characters while preserving useful punctuation."""
    if not text:
        return ""
    replacements = {
        "\u00a0": " ",
        "\u2010": "-",
        "\u2011": "-",
        "\u2012": "-",
        "\u2013": "-",
        "\u2014": "-",
        "\u2212": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\x00": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def flattened_text(text: str) -> str:
    return re.sub(r"\s+", " ", normalize_text(text)).strip()


def alnum_quality(text: str) -> int:
    return sum(ch.isalnum() for ch in text)


def should_ocr(embedded_text: str) -> bool:
    return alnum_quality(embedded_text) < MIN_EMBEDDED_TEXT_ALNUM


def unique_preserve(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        key = value.strip().lower()
        if value.strip() and key not in seen:
            seen.add(key)
            out.append(value.strip())
    return out


def make_snippet(text: str, start: int, end: int, radius: int = 190) -> str:
    flat = flattened_text(text)
    # Match indices were generated from flattened text in analyze_page, so use
    # them directly against the same representation.
    left = max(0, start - radius)
    right = min(len(flat), end + radius)
    snippet = flat[left:right]
    if left > 0:
        snippet = "..." + snippet
    if right < len(flat):
        snippet += "..."
    return snippet


def format_pressure(raw: str) -> str:
    value = raw.replace(" ", "")
    return f"{value} psi"


def format_gradient(raw_value: str) -> str:
    cleaned = raw_value.replace(",", ".")
    if cleaned.startswith("."):
        cleaned = "0" + cleaned
    try:
        value = float(cleaned)
    except ValueError:
        return f"{raw_value} psi/ft"

    if 0.1 <= value <= 2.0:
        return f"{value:g} psi/ft"

    # OCR frequently loses a leading decimal. Do not silently replace the raw
    # result; show a plausible candidate for human review.
    candidate: Optional[float] = None
    for divisor in (10.0, 100.0, 1000.0):
        test = value / divisor
        if 0.1 <= test <= 2.0:
            candidate = test
            break
    if candidate is not None:
        return f"{value:g} psi/ft [OCR may mean {candidate:g}]"
    return f"{value:g} psi/ft [check value]"


def find_pattern_values(text: str, patterns: Sequence[re.Pattern[str]]) -> list[str]:
    values: list[str] = []
    for pattern in patterns:
        for match in pattern.finditer(text):
            raw = match.groupdict().get("value")
            if raw:
                values.append(format_pressure(raw))
    return unique_preserve(values)


def extract_treating_values(text: str) -> list[str]:
    values: list[str] = []
    for pattern in TREATING_PATTERNS:
        for match in pattern.finditer(text):
            groups = match.groupdict()
            if groups.get("values"):
                nums = re.findall(NUMBER_VALUE, groups["values"])
                if nums:
                    values.append(" / ".join(nums) + " psi (Max/Min/Avg TP)")
            elif groups.get("value"):
                values.append(format_pressure(groups["value"]))
    return unique_preserve(values)


def extract_api_numbers(text: str, allow_unlabeled_8: bool = False) -> list[str]:
    """Extract California-style API numbers while limiting date false positives.

    Unlabeled 8-digit strings are accepted only for filenames. Within OCR text,
    an 8-digit number must be near an API label. Standalone 10-digit values that
    start with California's state code 04 are accepted anywhere.
    """
    candidates: list[str] = []

    def add(county: str, serial: str) -> None:
        county_value = int(county)
        if 1 <= county_value <= 199:
            candidates.append(f"04{county}{serial}")

    explicit = re.compile(
        r"(?:a\.?p\.?i\.?|api)\s*(?:no\.?|number|#)?\s*[:#-]?\s*"
        r"(?:04[\s-]?)?(?P<county>\d{3})[\s-]?(?P<serial>\d{5})(?!\d)",
        re.IGNORECASE,
    )
    for match in explicit.finditer(text):
        add(match.group("county"), match.group("serial"))

    ten_digit = re.compile(r"(?<!\d)04[\s-]?(?P<county>\d{3})[\s-]?(?P<serial>\d{5})(?!\d)")
    for match in ten_digit.finditer(text):
        add(match.group("county"), match.group("serial"))

    if allow_unlabeled_8:
        eight_digit = re.compile(r"(?<!\d)(?P<county>\d{3})[\s_-]?(?P<serial>\d{5})(?!\d)")
        for match in eight_digit.finditer(text):
            add(match.group("county"), match.group("serial"))

    return unique_preserve(candidates)


def extract_well_name(text: str, file_stem: str) -> str:
    normalized = normalize_text(text)
    lines = [re.sub(r"\s+", " ", line).strip() for line in normalized.splitlines()]

    line_patterns = (
        re.compile(r"^well\s*(?:name|no\.?|number|#)\s*[:\-]?\s*[\"']?(.{2,70})$", re.I),
        re.compile(r"^well\s*[:\-]\s*[\"']?(.{2,70})$", re.I),
        re.compile(r"^well\s+designation\s*[:\-]?\s*[\"']?(.{2,70})$", re.I),
    )
    for line in lines[:120]:
        for pattern in line_patterns:
            match = pattern.match(line)
            if match:
                candidate = match.group(1).strip(" \t\"'")
                candidate = re.split(r"\s+(?:sec\.?|section|field|county|api)\b", candidate, maxsplit=1, flags=re.I)[0]
                if 2 <= len(candidate) <= 70 and not candidate.lower().startswith(("data", "history")):
                    return candidate

    quoted = re.search(r"\bwell\s+[\"']([^\"']{2,70})[\"']", normalized, re.I)
    if quoted:
        return quoted.group(1).strip()

    # Conservative filename fallback. Do not invent a well name from generic
    # CalGEM file names such as 03022876_DATA_2019-06-21.
    fallback = re.sub(r"(?<!\d)(?:04)?\d{3}[\s_-]?\d{5}(?!\d)", "", file_stem)
    fallback = re.sub(r"[_\-]+", " ", fallback)
    fallback = re.sub(r"\b(?:data|well\s*file|calgem|step\s*rate\s*test|pdf)\b", "", fallback, flags=re.I)
    fallback = re.sub(r"\b\d{4}\s+\d{2}\s+\d{2}\b", "", fallback)
    fallback = re.sub(r"\s+", " ", fallback).strip(" -_")
    if 2 <= len(fallback) <= 70 and re.search(r"[A-Za-z]", fallback):
        return fallback
    return ""


# ------------------------------ OCR functions -------------------------------


def locate_tesseract(user_path: str = "") -> str:
    candidates: list[str] = []
    if user_path:
        candidates.append(user_path)

    env_path = os.environ.get("TESSERACT_CMD", "").strip()
    if env_path:
        candidates.append(env_path)

    saved_path_file = Path(__file__).resolve().parent / ".tesseract_path"
    try:
        saved_path = saved_path_file.read_text(encoding="utf-8").strip()
    except OSError:
        saved_path = ""
    if saved_path:
        candidates.append(saved_path)

    found = shutil.which("tesseract")
    if found:
        candidates.append(found)

    if os.name == "nt":
        program_files = os.environ.get("ProgramFiles", r"C:\Program Files")
        program_files_x86 = os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")
        local_app_data = os.environ.get("LOCALAPPDATA", "")
        candidates.extend(
            [
                str(Path(program_files) / "Tesseract-OCR" / "tesseract.exe"),
                str(Path(program_files_x86) / "Tesseract-OCR" / "tesseract.exe"),
            ]
        )
        if local_app_data:
            candidates.extend(
                [
                    str(Path(local_app_data) / "Programs" / "Tesseract-OCR" / "tesseract.exe"),
                    str(Path(local_app_data) / "Tesseract-OCR" / "tesseract.exe"),
                ]
            )

    seen: set[str] = set()
    for candidate in candidates:
        if not candidate:
            continue
        path = Path(candidate).expanduser()
        key = os.path.normcase(str(path))
        if key in seen:
            continue
        seen.add(key)
        if path.is_file():
            return str(path)
    return ""


def configure_tesseract(path: str) -> None:
    if pytesseract is None:
        raise RuntimeError("pytesseract is not installed. Run: pip install pytesseract")
    resolved = locate_tesseract(path)
    if not resolved:
        raise RuntimeError(
            "Tesseract OCR was not found. Install Tesseract or browse to tesseract.exe in the GUI."
        )
    pytesseract.pytesseract.tesseract_cmd = resolved


def ocr_page(page: "fitz.Page", dpi: int) -> str:
    if pytesseract is None:
        raise RuntimeError("pytesseract is not installed")
    scale = max(1.0, float(dpi) / 72.0)
    pix = page.get_pixmap(
        matrix=fitz.Matrix(scale, scale),
        colorspace=fitz.csGRAY,
        alpha=False,
    )
    temp_path = ""
    try:
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as temp:
            temp_path = temp.name
        pix.save(temp_path)
        # PSM 6 works well on typed regulatory forms and well-history pages.
        return pytesseract.image_to_string(
            temp_path,
            lang="eng",
            config="--oem 3 --psm 6 -c preserve_interword_spaces=1",
        )
    finally:
        if temp_path:
            try:
                os.remove(temp_path)
            except OSError:
                pass


# ------------------------------- Page analysis ------------------------------


def analyze_page(text: str, page_number: int, source: str) -> PageEvidence:
    flat = flattened_text(text)
    evidence = PageEvidence(
        page=page_number,
        source=source,
        text_quality=alnum_quality(flat),
    )
    if not flat:
        return evidence

    for rule in KEYWORD_RULES:
        matches = list(rule.pattern.finditer(flat))[:MAX_MATCHES_PER_RULE_PER_PAGE]
        if not matches:
            continue
        evidence.categories.add(rule.category)
        evidence.category_weights[rule.category] = max(
            evidence.category_weights.get(rule.category, 0), rule.weight
        )
        for match in matches:
            if len(evidence.details) >= MAX_DETAILS_PER_FILE:
                break
            evidence.details.append(
                MatchDetail(
                    page=page_number,
                    source=source,
                    category=rule.category,
                    matched_text=match.group(0),
                    snippet=make_snippet(flat, match.start(), match.end()),
                )
            )

    evidence.gradients = unique_preserve(
        format_gradient(match.group("value")) for match in GRADIENT_RE.finditer(flat)
    )
    evidence.isip_values = find_pattern_values(flat, ISIP_PATTERNS)
    evidence.closure_values = find_pattern_values(flat, CLOSURE_PATTERNS)
    evidence.breakdown_values = find_pattern_values(flat, BREAKDOWN_PATTERNS)
    evidence.treating_values = extract_treating_values(flat)
    evidence.rate_values = unique_preserve(
        f"{match.group('value')} {match.group('unit').upper().replace(' ', '')}"
        for match in RATE_RE.finditer(flat)
    )
    evidence.proppant_values = unique_preserve(
        "{} lb {}".format(
            match.group("amount"),
            re.sub(r"\s+", " ", match.group("type")).strip(),
        )
        for match in PROPPANT_RE.finditer(flat)
    )

    evidence.actual_marker = bool(ACTUAL_MARKER_RE.search(flat)) or bool(
        evidence.isip_values
        or evidence.closure_values
        or evidence.breakdown_values
        or evidence.treating_values
    )
    evidence.planned_marker = bool(PLANNED_MARKER_RE.search(flat))
    evidence.assumed_marker = bool(ASSUMED_MARKER_RE.search(flat))

    # Add synthetic detail rows for numeric values so the reviewer can find the
    # exact page even when OCR slightly damages the keyword itself.
    synthetic_values = (
        [("Fracture gradient", value) for value in evidence.gradients]
        + [("ISIP", value) for value in evidence.isip_values]
        + [("Closure pressure", value) for value in evidence.closure_values]
        + [("Breakdown pressure", value) for value in evidence.breakdown_values]
        + [("Treating pressure", value) for value in evidence.treating_values]
        + [("Rate data", value) for value in evidence.rate_values]
        + [("Proppant data", value) for value in evidence.proppant_values]
    )
    for category, value in synthetic_values:
        if len(evidence.details) >= MAX_DETAILS_PER_FILE:
            break
        evidence.details.append(
            MatchDetail(
                page=page_number,
                source=source,
                category=category,
                matched_text=value,
                snippet=flat[:700] + ("..." if len(flat) > 700 else ""),
            )
        )

    return evidence


def preliminary_strong_evidence(pages: Sequence[PageEvidence]) -> bool:
    categories = set().union(*(page.categories for page in pages)) if pages else set()
    gradients = any(page.gradients for page in pages)
    pressures = any(
        page.isip_values
        or page.closure_values
        or page.breakdown_values
        or page.treating_values
        for page in pages
    )
    rates = any(page.rate_values for page in pages)
    actual = any(page.actual_marker for page in pages)

    test_type = bool(categories & {"Step-rate test", "Mini-frac / DFIT", "Leak-off / FIT"})
    hydraulic = "Hydraulic-fracture treatment" in categories
    if test_type and (gradients or pressures):
        return True
    if hydraulic and actual and (gradients or pressures):
        return True
    if gradients and actual and pressures:
        return True
    return False


def classify_document(pages: Sequence[PageEvidence], image_only_not_ocrd: int) -> tuple[str, int, str]:
    if not pages:
        return "NO CLEAR DATA", 0, "No readable page text was available."

    category_weights: dict[str, int] = {}
    for page in pages:
        for category, weight in page.category_weights.items():
            category_weights[category] = max(category_weights.get(category, 0), weight)

    categories = set(category_weights)
    gradients = unique_preserve(v for p in pages for v in p.gradients)
    isip = unique_preserve(v for p in pages for v in p.isip_values)
    closure = unique_preserve(v for p in pages for v in p.closure_values)
    breakdown = unique_preserve(v for p in pages for v in p.breakdown_values)
    treating = unique_preserve(v for p in pages for v in p.treating_values)
    rates = unique_preserve(v for p in pages for v in p.rate_values)
    proppant = unique_preserve(v for p in pages for v in p.proppant_values)

    actual = any(p.actual_marker for p in pages)
    planned = any(p.planned_marker for p in pages)
    assumed = any(p.assumed_marker for p in pages)

    score = sum(category_weights.values())
    score += 5 if gradients else 0
    score += 4 if isip else 0
    score += 4 if closure else 0
    score += 4 if breakdown else 0
    score += 4 if treating else 0
    score += 2 if rates else 0
    score += 1 if proppant else 0
    score += 3 if actual else 0

    test_type = bool(categories & {"Step-rate test", "Mini-frac / DFIT", "Leak-off / FIT"})
    hydraulic = "Hydraulic-fracture treatment" in categories
    pressure_metric = bool(isip or closure or breakdown or treating)

    if test_type and (gradients or pressure_metric or rates):
        status = "YES - Test/treatment data"
        note = "A test type and associated numeric pressure/rate/gradient evidence were found."
    elif hydraulic and actual and (gradients or pressure_metric or (rates and proppant)):
        status = "YES - Test/treatment data"
        note = "An actual hydraulic-fracture treatment record with numeric evidence was found."
    elif gradients and actual and pressure_metric:
        status = "YES - Test/treatment data"
        note = "Measured fracture-gradient and pressure evidence were found."
    elif (hydraulic or test_type) and planned and not actual:
        status = "POSSIBLE - Planned only"
        note = "The file appears to contain a proposed or planned test/treatment, but no clear execution record."
    elif gradients and assumed and not (test_type or hydraulic or actual):
        status = "REFERENCE/ASSUMED ONLY"
        note = "A fracture-gradient value appears to be assumed, estimated, permitted, or used as a reference."
    elif (
        score >= 7
        and (
            categories
            & {
                "Step-rate test",
                "Mini-frac / DFIT",
                "Leak-off / FIT",
                "Fracture gradient",
                "Breakdown pressure",
                "Closure pressure",
                "ISIP",
                "Hydraulic-fracture treatment",
                "Treating pressure",
            }
            or gradients
            or pressure_metric
            or (rates and proppant)
        )
    ):
        status = "POSSIBLE - Review"
        note = "Relevant keywords or values were found, but the evidence is not strong enough for an automatic YES."
    else:
        status = "NO CLEAR DATA"
        note = "No clear fracture-test or fracture-treatment evidence was found."

    if image_only_not_ocrd:
        note += f" {image_only_not_ocrd} image-only page(s) were not OCR'd, so this result may be incomplete."
    return status, score, note


# ------------------------------- PDF scanning -------------------------------


def scan_pdf(
    pdf_path: Path,
    options: ScanOptions,
    stop_event: Optional[threading.Event] = None,
    page_callback: Optional[Callable[[int, int, str], None]] = None,
) -> ScanResult:
    start_time = time.time()
    if fitz is None:
        raise RuntimeError("PyMuPDF is not installed. Run: pip install pymupdf")

    if options.use_ocr:
        configure_tesseract(options.tesseract_path)

    page_evidence: list[PageEvidence] = []
    metadata_text_parts: list[str] = [pdf_path.name]
    ocr_pages = 0
    image_only_not_ocrd = 0
    pages_scanned = 0

    try:
        doc = fitz.open(str(pdf_path))
    except Exception as exc:
        return ScanResult(
            file_path=str(pdf_path),
            file_name=pdf_path.name,
            status="ERROR",
            score=0,
            api_numbers="; ".join(extract_api_numbers(pdf_path.name, allow_unlabeled_8=True)),
            well_name=extract_well_name("", pdf_path.stem),
            categories="",
            fracture_gradients="",
            isip_values="",
            closure_pressures="",
            breakdown_pressures="",
            treating_pressures="",
            rate_values="",
            proppant_values="",
            matched_pages="",
            pages_scanned=0,
            total_pages=0,
            ocr_pages=0,
            image_only_pages_not_ocrd=0,
            notes=f"Could not open PDF: {exc}",
            details=[],
            elapsed_seconds=time.time() - start_time,
        )

    total_pages = doc.page_count
    try:
        if getattr(doc, "needs_pass", False):
            raise RuntimeError("PDF is password-protected")

        for page_index in range(total_pages):
            if stop_event and stop_event.is_set():
                break

            page_number = page_index + 1
            if page_callback:
                page_callback(page_number, total_pages, pdf_path.name)

            page = doc.load_page(page_index)
            embedded = normalize_text(page.get_text("text", sort=True))
            text = embedded
            source = "embedded text"

            if should_ocr(embedded):
                if options.use_ocr:
                    try:
                        ocr_text = normalize_text(ocr_page(page, options.ocr_dpi))
                        ocr_pages += 1
                        # Prefer OCR if it contains more alphanumeric content.
                        if alnum_quality(ocr_text) > alnum_quality(embedded):
                            text = ocr_text
                            source = "OCR"
                    except Exception as exc:
                        image_only_not_ocrd += 1
                        source = f"OCR error: {exc}"
                else:
                    image_only_not_ocrd += 1
                    source = "image-only; OCR disabled"

            pages_scanned += 1
            evidence = analyze_page(text, page_number, source)
            page_evidence.append(evidence)

            # The first pages usually contain API and well name. Matched pages
            # are also useful metadata sources.
            if page_index < 3 or evidence.categories or evidence.gradients or evidence.isip_values:
                metadata_text_parts.append(text[:8000])

            if options.stop_after_strong_evidence and preliminary_strong_evidence(page_evidence):
                break
    except Exception as exc:
        details = [
            MatchDetail(
                page=0,
                source="scanner",
                category="Error",
                matched_text=str(exc),
                snippet=traceback.format_exc(limit=4),
            )
        ]
        return ScanResult(
            file_path=str(pdf_path),
            file_name=pdf_path.name,
            status="ERROR",
            score=0,
            api_numbers="; ".join(extract_api_numbers("\n".join(metadata_text_parts))),
            well_name=extract_well_name("\n".join(metadata_text_parts), pdf_path.stem),
            categories="",
            fracture_gradients="",
            isip_values="",
            closure_pressures="",
            breakdown_pressures="",
            treating_pressures="",
            rate_values="",
            proppant_values="",
            matched_pages="",
            pages_scanned=pages_scanned,
            total_pages=total_pages,
            ocr_pages=ocr_pages,
            image_only_pages_not_ocrd=image_only_not_ocrd,
            notes=f"Scan error: {exc}",
            details=details,
            elapsed_seconds=time.time() - start_time,
        )
    finally:
        doc.close()

    status, score, notes = classify_document(page_evidence, image_only_not_ocrd)

    categories = sorted(set().union(*(p.categories for p in page_evidence)) if page_evidence else set())
    gradients = unique_preserve(v for p in page_evidence for v in p.gradients)
    isip = unique_preserve(v for p in page_evidence for v in p.isip_values)
    closure = unique_preserve(v for p in page_evidence for v in p.closure_values)
    breakdown = unique_preserve(v for p in page_evidence for v in p.breakdown_values)
    treating = unique_preserve(v for p in page_evidence for v in p.treating_values)
    rates = unique_preserve(v for p in page_evidence for v in p.rate_values)
    proppant = unique_preserve(v for p in page_evidence for v in p.proppant_values)
    details = [detail for p in page_evidence for detail in p.details][:MAX_DETAILS_PER_FILE]
    matched_pages = sorted(
        {
            p.page
            for p in page_evidence
            if p.categories
            or p.gradients
            or p.isip_values
            or p.closure_values
            or p.breakdown_values
            or p.treating_values
        }
    )

    metadata_text = "\n".join(metadata_text_parts)
    api_numbers = unique_preserve(
        extract_api_numbers(pdf_path.name, allow_unlabeled_8=True)
        + extract_api_numbers(metadata_text)
    )
    well_name = extract_well_name(metadata_text, pdf_path.stem)

    if pages_scanned < total_pages and options.stop_after_strong_evidence:
        notes += f" Scan stopped after page {pages_scanned} because strong evidence was found."

    return ScanResult(
        file_path=str(pdf_path),
        file_name=pdf_path.name,
        status=status,
        score=score,
        api_numbers="; ".join(api_numbers),
        well_name=well_name,
        categories="; ".join(categories),
        fracture_gradients="; ".join(gradients),
        isip_values="; ".join(isip),
        closure_pressures="; ".join(closure),
        breakdown_pressures="; ".join(breakdown),
        treating_pressures="; ".join(treating),
        rate_values="; ".join(rates[:20]),
        proppant_values="; ".join(proppant[:20]),
        matched_pages=", ".join(str(page) for page in matched_pages),
        pages_scanned=pages_scanned,
        total_pages=total_pages,
        ocr_pages=ocr_pages,
        image_only_pages_not_ocrd=image_only_not_ocrd,
        notes=notes,
        details=details,
        elapsed_seconds=time.time() - start_time,
    )


def discover_pdfs(folder: Path, recursive: bool) -> list[Path]:
    iterator = folder.rglob("*.pdf") if recursive else folder.glob("*.pdf")
    return sorted((p for p in iterator if p.is_file()), key=lambda p: str(p).lower())


# ------------------------------- Export logic -------------------------------


def aggregate_wells(results: Sequence[ScanResult]) -> list[dict[str, object]]:
    groups: dict[str, list[ScanResult]] = {}
    for result in results:
        api = result.api_numbers.split(";")[0].strip() if result.api_numbers else ""
        key = api or result.well_name.strip().lower() or f"FILE::{result.file_name.lower()}"
        groups.setdefault(key, []).append(result)

    rows: list[dict[str, object]] = []
    for key, items in groups.items():
        best = max(items, key=lambda r: STATUS_RANK.get(r.status, -1))
        rows.append(
            {
                "API Number": best.api_numbers,
                "Well Name": best.well_name,
                "Best Status": best.status,
                "Best Score": max(item.score for item in items),
                "Categories": "; ".join(unique_preserve(v for i in items for v in i.categories.split(";") if v.strip())),
                "Fracture Gradient": "; ".join(unique_preserve(v for i in items for v in i.fracture_gradients.split(";") if v.strip())),
                "ISIP": "; ".join(unique_preserve(v for i in items for v in i.isip_values.split(";") if v.strip())),
                "Closure Pressure": "; ".join(unique_preserve(v for i in items for v in i.closure_pressures.split(";") if v.strip())),
                "Breakdown Pressure": "; ".join(unique_preserve(v for i in items for v in i.breakdown_pressures.split(";") if v.strip())),
                "Treating Pressure": "; ".join(unique_preserve(v for i in items for v in i.treating_pressures.split(";") if v.strip())),
                "Supporting PDF Count": len(items),
                "Supporting Files": "; ".join(item.file_name for item in items),
                "Supporting Paths": "; ".join(item.file_path for item in items),
            }
        )
    return sorted(rows, key=lambda row: (str(row["API Number"]), str(row["Well Name"])))


def _safe_excel_value(value: object) -> object:
    if isinstance(value, str) and len(value) > MAX_EXCEL_CELL:
        return value[:MAX_EXCEL_CELL] + "..."
    return value


def export_xlsx(results: Sequence[ScanResult], output_path: Path, settings: Optional[dict[str, object]] = None) -> None:
    if Workbook is None:
        raise RuntimeError("openpyxl is not installed. Run: pip install openpyxl")

    workbook = Workbook()
    summary = workbook.active
    summary.title = "PDF Summary"

    summary_rows = [result.summary_row() for result in results]
    headers = list(summary_rows[0].keys()) if summary_rows else list(
        ScanResult(
            "", "", "", 0, "", "", "", "", "", "", "", "", "", "", "", 0, 0, 0, 0, ""
        ).summary_row().keys()
    )
    summary.append(headers)
    for row in summary_rows:
        summary.append([_safe_excel_value(row.get(header, "")) for header in headers])

    details_ws = workbook.create_sheet("Match Details")
    detail_headers = ["File Name", "File Path", "Status", "API Number(s)", "Well Name", "Page", "Source", "Category", "Matched Text", "Snippet"]
    details_ws.append(detail_headers)
    for result in results:
        for detail in result.details:
            details_ws.append(
                [
                    result.file_name,
                    result.file_path,
                    result.status,
                    result.api_numbers,
                    result.well_name,
                    detail.page,
                    detail.source,
                    detail.category,
                    detail.matched_text,
                    _safe_excel_value(detail.snippet),
                ]
            )

    well_ws = workbook.create_sheet("Well Summary")
    well_rows = aggregate_wells(results)
    well_headers = list(well_rows[0].keys()) if well_rows else [
        "API Number", "Well Name", "Best Status", "Best Score", "Categories", "Fracture Gradient",
        "ISIP", "Closure Pressure", "Breakdown Pressure", "Treating Pressure",
        "Supporting PDF Count", "Supporting Files", "Supporting Paths"
    ]
    well_ws.append(well_headers)
    for row in well_rows:
        well_ws.append([_safe_excel_value(row.get(header, "")) for header in well_headers])

    settings_ws = workbook.create_sheet("Settings")
    settings_ws.append(["Setting", "Value"])
    for key, value in (settings or {}).items():
        settings_ws.append([key, str(value)])
    settings_ws.append(["Application", f"{APP_TITLE} v{APP_VERSION}"])
    settings_ws.append(["Export Time", time.strftime("%Y-%m-%d %H:%M:%S")])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    status_fills = {
        "YES - Test/treatment data": PatternFill("solid", fgColor="C6EFCE"),
        "POSSIBLE - Review": PatternFill("solid", fgColor="FFF2CC"),
        "POSSIBLE - Planned only": PatternFill("solid", fgColor="FCE4D6"),
        "REFERENCE/ASSUMED ONLY": PatternFill("solid", fgColor="DDEBF7"),
        "NO CLEAR DATA": PatternFill("solid", fgColor="E7E6E6"),
        "ERROR": PatternFill("solid", fgColor="F4CCCC"),
    }

    for ws in (summary, details_ws, well_ws, settings_ws):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Color status cells.
    if summary.max_row >= 2:
        status_col = headers.index("Status") + 1
        path_col = headers.index("File Path") + 1
        for row in range(2, summary.max_row + 1):
            status = str(summary.cell(row, status_col).value or "")
            summary.cell(row, status_col).fill = status_fills.get(status, PatternFill())
            path_cell = summary.cell(row, path_col)
            path = str(path_cell.value or "")
            if path:
                path_cell.hyperlink = Path(path).resolve().as_uri()
                path_cell.style = "Hyperlink"

    if well_ws.max_row >= 2 and "Best Status" in well_headers:
        status_col = well_headers.index("Best Status") + 1
        for row in range(2, well_ws.max_row + 1):
            status = str(well_ws.cell(row, status_col).value or "")
            well_ws.cell(row, status_col).fill = status_fills.get(status, PatternFill())

    # Practical column widths.
    widths = {
        "Status": 26,
        "Score": 9,
        "API Number(s)": 18,
        "API Number": 18,
        "Well Name": 28,
        "Categories": 38,
        "Fracture Gradient": 26,
        "ISIP": 20,
        "Closure Pressure": 20,
        "Breakdown Pressure": 20,
        "Treating Pressure": 28,
        "Rate Data": 28,
        "Proppant Data": 28,
        "Matched Pages": 15,
        "Notes": 55,
        "File Name": 40,
        "File Path": 75,
        "Snippet": 100,
        "Supporting Files": 70,
        "Supporting Paths": 100,
    }
    for ws in (summary, details_ws, well_ws, settings_ws):
        for col_idx, header_cell in enumerate(ws[1], start=1):
            header = str(header_cell.value or "")
            default_width = min(max(len(header) + 2, 12), 30)
            ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(header, default_width)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def export_csv(results: Sequence[ScanResult], output_path: Path) -> None:
    rows = [result.summary_row() for result in results]
    headers = list(rows[0].keys()) if rows else []
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def export_results(results: Sequence[ScanResult], output_path: Path, settings: Optional[dict[str, object]] = None) -> None:
    if output_path.suffix.lower() == ".csv":
        export_csv(results, output_path)
    else:
        if output_path.suffix.lower() != ".xlsx":
            output_path = output_path.with_suffix(".xlsx")
        export_xlsx(results, output_path, settings)


# --------------------------------- GUI app ----------------------------------


class FractureScannerApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title(f"{APP_TITLE} v{APP_VERSION}")
        self.root.geometry("1500x860")
        self.root.minsize(1100, 700)

        self.message_queue: queue.Queue[tuple] = queue.Queue()
        self.stop_event = threading.Event()
        self.scan_thread: Optional[threading.Thread] = None
        self.results: list[ScanResult] = []
        self.result_by_item: dict[str, ScanResult] = {}

        detected_tesseract = locate_tesseract()
        self.folder_var = tk.StringVar()
        self.recursive_var = tk.BooleanVar(value=True)
        self.ocr_var = tk.BooleanVar(value=bool(detected_tesseract))
        self.early_stop_var = tk.BooleanVar(value=True)
        self.dpi_var = tk.StringVar(value=str(DEFAULT_OCR_DPI))
        self.workers_var = tk.StringVar(value="2")
        self.tesseract_var = tk.StringVar(value=detected_tesseract)
        initial_status = "Select a folder containing PDF well files."
        if not detected_tesseract:
            initial_status += " Tesseract was not found, so OCR is currently disabled."
        self.status_var = tk.StringVar(value=initial_status)
        self.progress_text_var = tk.StringVar(value="0 / 0")

        self._configure_style()
        self._build_ui()
        self.root.after(150, self._poll_queue)

    def _configure_style(self) -> None:
        default_font = tkfont.nametofont("TkDefaultFont")
        default_font.configure(size=10)
        tkfont.nametofont("TkTextFont").configure(size=10)
        tkfont.nametofont("TkHeadingFont").configure(size=11, weight="bold")

        style = ttk.Style(self.root)
        try:
            style.theme_use("vista" if os.name == "nt" else "clam")
        except tk.TclError:
            pass
        style.configure("Treeview", rowheight=27, font=("Segoe UI", 10))
        style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
        style.configure("TButton", padding=(10, 6))
        style.configure("TLabelframe.Label", font=("Segoe UI", 10, "bold"))

    def _build_ui(self) -> None:
        main = ttk.Frame(self.root, padding=10)
        main.pack(fill="both", expand=True)

        source_frame = ttk.LabelFrame(main, text="PDF source and scan options", padding=10)
        source_frame.pack(fill="x")
        source_frame.columnconfigure(1, weight=1)

        ttk.Label(source_frame, text="PDF folder:").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(source_frame, textvariable=self.folder_var).grid(row=0, column=1, sticky="ew", pady=4)
        ttk.Button(source_frame, text="Browse...", command=self._browse_folder).grid(row=0, column=2, padx=(8, 0), pady=4)

        ttk.Checkbutton(source_frame, text="Include subfolders", variable=self.recursive_var).grid(row=1, column=0, sticky="w", pady=4)
        ttk.Checkbutton(source_frame, text="OCR image-only pages", variable=self.ocr_var).grid(row=1, column=1, sticky="w", pady=4)
        ttk.Checkbutton(source_frame, text="Stop a PDF after strong evidence is found", variable=self.early_stop_var).grid(row=1, column=2, sticky="w", padx=(8, 0), pady=4)

        ttk.Label(source_frame, text="OCR DPI:").grid(row=2, column=0, sticky="w", pady=4)
        dpi_combo = ttk.Combobox(source_frame, textvariable=self.dpi_var, width=8, state="readonly", values=(120, 150, 180, 200, 250, 300))
        dpi_combo.grid(row=2, column=1, sticky="w", pady=4)

        worker_frame = ttk.Frame(source_frame)
        worker_frame.grid(row=2, column=2, sticky="w", padx=(8, 0), pady=4)
        ttk.Label(worker_frame, text="Parallel PDF workers:").pack(side="left")
        ttk.Combobox(
            worker_frame,
            textvariable=self.workers_var,
            width=4,
            state="readonly",
            values=(1, 2, 3, 4),
        ).pack(side="left", padx=(6, 0))

        ttk.Label(source_frame, text="Tesseract executable:").grid(row=3, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(source_frame, textvariable=self.tesseract_var).grid(row=3, column=1, sticky="ew", pady=4)
        ttk.Button(source_frame, text="Browse...", command=self._browse_tesseract).grid(row=3, column=2, padx=(8, 0), pady=4)

        button_frame = ttk.Frame(main, padding=(0, 10, 0, 8))
        button_frame.pack(fill="x")
        self.start_button = ttk.Button(button_frame, text="Start Scan", command=self._start_scan)
        self.start_button.pack(side="left")
        self.stop_button = ttk.Button(button_frame, text="Stop", command=self._stop_scan, state="disabled")
        self.stop_button.pack(side="left", padx=8)
        self.export_button = ttk.Button(button_frame, text="Export Results...", command=self._export_results, state="disabled")
        self.export_button.pack(side="left")
        self.open_button = ttk.Button(button_frame, text="Open Selected PDF", command=self._open_selected, state="disabled")
        self.open_button.pack(side="left", padx=8)
        self.clear_button = ttk.Button(button_frame, text="Clear", command=self._clear_results)
        self.clear_button.pack(side="left")

        ttk.Label(button_frame, textvariable=self.progress_text_var).pack(side="right")

        self.progress = ttk.Progressbar(main, mode="determinate")
        self.progress.pack(fill="x", pady=(0, 6))
        ttk.Label(main, textvariable=self.status_var).pack(fill="x", pady=(0, 8))

        pane = ttk.Panedwindow(main, orient="vertical")
        pane.pack(fill="both", expand=True)

        table_frame = ttk.Frame(pane)
        details_frame = ttk.LabelFrame(pane, text="Selected PDF evidence", padding=6)
        pane.add(table_frame, weight=3)
        pane.add(details_frame, weight=2)

        columns = (
            "status", "score", "api", "well", "categories", "gradient", "isip", "pages", "ocr", "file"
        )
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings", selectmode="browse")
        headings = {
            "status": "Status",
            "score": "Score",
            "api": "API Number(s)",
            "well": "Well Name",
            "categories": "Categories",
            "gradient": "Fracture Gradient",
            "isip": "ISIP",
            "pages": "Matched Pages",
            "ocr": "OCR Pages",
            "file": "PDF File",
        }
        widths = {
            "status": 205,
            "score": 65,
            "api": 140,
            "well": 180,
            "categories": 280,
            "gradient": 190,
            "isip": 140,
            "pages": 100,
            "ocr": 85,
            "file": 320,
        }
        for column in columns:
            self.tree.heading(column, text=headings[column])
            anchor = "center" if column in {"score", "pages", "ocr"} else "w"
            self.tree.column(column, width=widths[column], minwidth=55, anchor=anchor, stretch=column in {"categories", "file"})

        self.tree.tag_configure("yes", background="#d9ead3")
        self.tree.tag_configure("possible", background="#fff2cc")
        self.tree.tag_configure("planned", background="#fce5cd")
        self.tree.tag_configure("reference", background="#d9eaf7")
        self.tree.tag_configure("no", background="#eeeeee")
        self.tree.tag_configure("error", background="#f4cccc")

        y_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        x_scroll = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll.grid(row=1, column=0, sticky="ew")
        table_frame.rowconfigure(0, weight=1)
        table_frame.columnconfigure(0, weight=1)

        self.tree.bind("<<TreeviewSelect>>", self._show_selected_details)
        self.tree.bind("<Double-1>", lambda _event: self._open_selected())

        self.details_text = tk.Text(details_frame, wrap="word", font=("Consolas", 10), padx=8, pady=8)
        details_scroll = ttk.Scrollbar(details_frame, orient="vertical", command=self.details_text.yview)
        self.details_text.configure(yscrollcommand=details_scroll.set, state="disabled")
        self.details_text.pack(side="left", fill="both", expand=True)
        details_scroll.pack(side="right", fill="y")

    def _browse_folder(self) -> None:
        selected = filedialog.askdirectory(title="Select folder containing PDF files")
        if selected:
            self.folder_var.set(selected)

    def _browse_tesseract(self) -> None:
        selected = filedialog.askopenfilename(
            title="Select Tesseract executable",
            filetypes=[("Tesseract executable", "tesseract.exe" if os.name == "nt" else "tesseract"), ("All files", "*")],
        )
        if selected:
            self.tesseract_var.set(selected)

    def _validate_options(self) -> Optional[tuple[Path, ScanOptions, int]]:
        folder = Path(self.folder_var.get().strip())
        if not folder.is_dir():
            messagebox.showerror(APP_TITLE, "Please select a valid PDF folder.")
            return None
        if fitz is None:
            messagebox.showerror(APP_TITLE, "PyMuPDF is missing from the managed environment. Run Repair_or_Update_Environment.bat.")
            return None
        try:
            dpi = int(self.dpi_var.get())
            workers = int(self.workers_var.get())
        except ValueError:
            messagebox.showerror(APP_TITLE, "OCR DPI and worker count must be whole numbers.")
            return None
        workers = max(1, min(workers, 4))
        if self.ocr_var.get():
            if pytesseract is None:
                messagebox.showerror(APP_TITLE, "pytesseract is missing from the managed environment. Run Repair_or_Update_Environment.bat.")
                return None
            tess = locate_tesseract(self.tesseract_var.get().strip())
            if not tess:
                messagebox.showerror(APP_TITLE, "Tesseract OCR was not found. Browse to tesseract.exe or disable OCR.")
                return None
            self.tesseract_var.set(tess)

        options = ScanOptions(
            use_ocr=self.ocr_var.get(),
            ocr_dpi=dpi,
            tesseract_path=self.tesseract_var.get().strip(),
            stop_after_strong_evidence=self.early_stop_var.get(),
        )
        return folder, options, workers

    def _start_scan(self) -> None:
        validated = self._validate_options()
        if validated is None:
            return
        folder, options, workers = validated
        pdfs = discover_pdfs(folder, self.recursive_var.get())
        if not pdfs:
            messagebox.showinfo(APP_TITLE, "No PDF files were found in the selected folder.")
            return

        self._clear_results(confirm=False)
        self.stop_event.clear()
        self.start_button.configure(state="disabled")
        self.stop_button.configure(state="normal")
        self.export_button.configure(state="disabled")
        self.open_button.configure(state="disabled")
        self.clear_button.configure(state="disabled")
        self.progress.configure(maximum=len(pdfs), value=0)
        self.progress_text_var.set(f"0 / {len(pdfs)}")
        self.status_var.set(f"Preparing to scan {len(pdfs)} PDF file(s)...")

        self.scan_thread = threading.Thread(
            target=self._scan_worker,
            args=(pdfs, options, workers),
            daemon=True,
        )
        self.scan_thread.start()

    def _scan_worker(self, pdfs: Sequence[Path], options: ScanOptions, workers: int) -> None:
        total = len(pdfs)

        def scan_one(source_index: int, pdf_path: Path) -> tuple[int, ScanResult]:
            def page_callback(page: int, total_pages: int, name: str) -> None:
                self.message_queue.put(("page", source_index, total, page, total_pages, name))

            try:
                result = scan_pdf(pdf_path, options, self.stop_event, page_callback)
            except Exception as exc:
                result = ScanResult(
                    file_path=str(pdf_path),
                    file_name=pdf_path.name,
                    status="ERROR",
                    score=0,
                    api_numbers="; ".join(extract_api_numbers(pdf_path.name, allow_unlabeled_8=True)),
                    well_name=extract_well_name("", pdf_path.stem),
                    categories="",
                    fracture_gradients="",
                    isip_values="",
                    closure_pressures="",
                    breakdown_pressures="",
                    treating_pressures="",
                    rate_values="",
                    proppant_values="",
                    matched_pages="",
                    pages_scanned=0,
                    total_pages=0,
                    ocr_pages=0,
                    image_only_pages_not_ocrd=0,
                    notes=f"Unexpected scan failure: {exc}",
                    details=[MatchDetail(0, "scanner", "Error", str(exc), traceback.format_exc(limit=5))],
                    elapsed_seconds=0.0,
                )
            return source_index, result

        completed = 0
        with concurrent.futures.ThreadPoolExecutor(max_workers=workers, thread_name_prefix="fracture-scan") as executor:
            future_map = {
                executor.submit(scan_one, index, pdf_path): (index, pdf_path)
                for index, pdf_path in enumerate(pdfs, start=1)
            }
            for future in concurrent.futures.as_completed(future_map):
                try:
                    source_index, result = future.result()
                except Exception as exc:
                    source_index, pdf_path = future_map[future]
                    result = ScanResult(
                        file_path=str(pdf_path),
                        file_name=pdf_path.name,
                        status="ERROR",
                        score=0,
                        api_numbers="; ".join(extract_api_numbers(pdf_path.name, allow_unlabeled_8=True)),
                        well_name=extract_well_name("", pdf_path.stem),
                        categories="",
                        fracture_gradients="",
                        isip_values="",
                        closure_pressures="",
                        breakdown_pressures="",
                        treating_pressures="",
                        rate_values="",
                        proppant_values="",
                        matched_pages="",
                        pages_scanned=0,
                        total_pages=0,
                        ocr_pages=0,
                        image_only_pages_not_ocrd=0,
                        notes=f"Worker failure: {exc}",
                        details=[MatchDetail(0, "scanner", "Error", str(exc), traceback.format_exc(limit=5))],
                        elapsed_seconds=0.0,
                    )
                completed += 1
                self.message_queue.put(("result", completed, total, source_index, result))

                if self.stop_event.is_set():
                    for pending in future_map:
                        pending.cancel()
                    break

        self.message_queue.put(("done", self.stop_event.is_set(), completed, total))

    def _stop_scan(self) -> None:
        self.stop_event.set()
        self.status_var.set("Stopping after the current page...")
        self.stop_button.configure(state="disabled")

    def _poll_queue(self) -> None:
        try:
            while True:
                message = self.message_queue.get_nowait()
                kind = message[0]
                if kind == "page":
                    _, file_index, file_total, page, page_total, name = message
                    self.status_var.set(f"Scanning file {file_index}/{file_total}: {name} - page {page}/{page_total}")
                elif kind == "result":
                    _, completed, total, _source_index, result = message
                    self.results.append(result)
                    self._insert_result(result)
                    self.progress.configure(value=completed)
                    self.progress_text_var.set(f"{completed} / {total}")
                    self.status_var.set(f"Finished {result.file_name}: {result.status}")
                elif kind == "done":
                    _, stopped, _result_count_at_worker, total = message
                    self.start_button.configure(state="normal")
                    self.stop_button.configure(state="disabled")
                    self.clear_button.configure(state="normal")
                    if self.results:
                        self.export_button.configure(state="normal")
                    self.status_var.set(
                        f"Scan {'stopped' if stopped else 'complete'}: {len(self.results)} of {total} PDF file(s) processed."
                    )
        except queue.Empty:
            pass
        self.root.after(150, self._poll_queue)

    def _insert_result(self, result: ScanResult) -> None:
        tag = {
            "YES - Test/treatment data": "yes",
            "POSSIBLE - Review": "possible",
            "POSSIBLE - Planned only": "planned",
            "REFERENCE/ASSUMED ONLY": "reference",
            "NO CLEAR DATA": "no",
            "ERROR": "error",
        }.get(result.status, "no")
        item = self.tree.insert(
            "",
            "end",
            values=(
                result.status,
                result.score,
                result.api_numbers,
                result.well_name,
                result.categories,
                result.fracture_gradients,
                result.isip_values,
                result.matched_pages,
                result.ocr_pages,
                result.file_name,
            ),
            tags=(tag,),
        )
        self.result_by_item[item] = result

    def _show_selected_details(self, _event: object = None) -> None:
        selection = self.tree.selection()
        if not selection:
            self.open_button.configure(state="disabled")
            return
        result = self.result_by_item.get(selection[0])
        if result is None:
            return
        self.open_button.configure(state="normal")

        lines = [
            f"FILE: {result.file_path}",
            f"STATUS: {result.status}",
            f"SCORE: {result.score}",
            f"API: {result.api_numbers or '(not found)'}",
            f"WELL: {result.well_name or '(not found)'}",
            f"CATEGORIES: {result.categories or '(none)'}",
            f"FRACTURE GRADIENT: {result.fracture_gradients or '(none)'}",
            f"ISIP: {result.isip_values or '(none)'}",
            f"CLOSURE PRESSURE: {result.closure_pressures or '(none)'}",
            f"BREAKDOWN PRESSURE: {result.breakdown_pressures or '(none)'}",
            f"TREATING PRESSURE: {result.treating_pressures or '(none)'}",
            f"RATE DATA: {result.rate_values or '(none)'}",
            f"PROPPANT DATA: {result.proppant_values or '(none)'}",
            f"MATCHED PAGES: {result.matched_pages or '(none)'}",
            f"PAGES SCANNED: {result.pages_scanned}/{result.total_pages}; OCR PAGES: {result.ocr_pages}",
            f"NOTES: {result.notes}",
            "",
            "MATCH DETAILS",
            "=" * 90,
        ]
        for detail in result.details:
            lines.extend(
                [
                    f"Page {detail.page} | {detail.category} | {detail.source}",
                    f"Matched: {detail.matched_text}",
                    detail.snippet,
                    "-" * 90,
                ]
            )
        if not result.details:
            lines.append("No keyword snippets were captured.")

        self.details_text.configure(state="normal")
        self.details_text.delete("1.0", "end")
        self.details_text.insert("1.0", "\n".join(lines))
        self.details_text.configure(state="disabled")

    def _open_selected(self) -> None:
        selection = self.tree.selection()
        if not selection:
            return
        result = self.result_by_item.get(selection[0])
        if result is None:
            return
        path = Path(result.file_path)
        if not path.exists():
            messagebox.showerror(APP_TITLE, f"The file no longer exists:\n{path}")
            return
        try:
            if os.name == "nt":
                os.startfile(str(path))  # type: ignore[attr-defined]
            elif sys.platform == "darwin":
                subprocess.Popen(["open", str(path)])
            else:
                subprocess.Popen(["xdg-open", str(path)])
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"Could not open the PDF:\n{exc}")

    def _export_results(self) -> None:
        if not self.results:
            return
        initial_dir = self.folder_var.get().strip() or os.getcwd()
        selected = filedialog.asksaveasfilename(
            title="Export fracture scan results",
            initialdir=initial_dir,
            initialfile="fracture_data_scan_results.xlsx",
            defaultextension=".xlsx",
            filetypes=[("Excel workbook", "*.xlsx"), ("CSV summary", "*.csv")],
        )
        if not selected:
            return
        settings = {
            "Source Folder": self.folder_var.get(),
            "Recursive": self.recursive_var.get(),
            "OCR Enabled": self.ocr_var.get(),
            "OCR DPI": self.dpi_var.get(),
            "Parallel PDF Workers": self.workers_var.get(),
            "Tesseract": self.tesseract_var.get(),
            "Stop After Strong Evidence": self.early_stop_var.get(),
            "Application Version": APP_VERSION,
            "Python Executable": sys.executable,
            "Python Prefix": sys.prefix,
        }
        try:
            export_results(self.results, Path(selected), settings)
            messagebox.showinfo(APP_TITLE, f"Results exported successfully:\n{selected}")
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"Could not export results:\n{exc}")

    def _clear_results(self, confirm: bool = True) -> None:
        if confirm and self.results and not messagebox.askyesno(APP_TITLE, "Clear all current scan results?"):
            return
        self.results.clear()
        self.result_by_item.clear()
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.details_text.configure(state="normal")
        self.details_text.delete("1.0", "end")
        self.details_text.configure(state="disabled")
        self.progress.configure(value=0)
        self.progress_text_var.set("0 / 0")
        self.export_button.configure(state="disabled")
        self.open_button.configure(state="disabled")


# ------------------------ Managed environment startup -----------------------


def _managed_venv_dir() -> Path:
    return Path(__file__).resolve().parent / ".venv"


def _running_in_managed_venv() -> bool:
    try:
        return Path(sys.prefix).resolve() == _managed_venv_dir().resolve()
    except OSError:
        return False


def _base_python_for_bootstrap() -> str:
    executable = Path(sys.executable)
    if os.name == "nt" and executable.name.lower() == "pythonw.exe":
        console_python = executable.with_name("python.exe")
        if console_python.is_file():
            return str(console_python)
    return str(executable)


def relaunch_through_managed_environment() -> Optional[int]:
    """Relaunch through the package bootstrap when this file is run directly.

    The batch launcher already starts the app from .venv. This fallback makes
    direct execution of the .py file follow the same managed-environment path.
    """

    if os.environ.get("FRACTURE_SCANNER_SKIP_BOOTSTRAP") == "1":
        return None
    if os.environ.get("FRACTURE_SCANNER_BOOTSTRAPPED") == "1":
        return None
    if _running_in_managed_venv():
        return None

    bootstrap = Path(__file__).resolve().parent / "bootstrap_fracture_scanner.py"
    if not bootstrap.is_file():
        return None

    command = [_base_python_for_bootstrap(), str(bootstrap), "--", *sys.argv[1:]]
    try:
        return subprocess.call(command, cwd=str(bootstrap.parent))
    except OSError as exc:
        print(f"ERROR: could not start the managed-environment bootstrap: {exc}", file=sys.stderr)
        return 2


# -------------------------------- CLI mode ----------------------------------


def run_cli(args: argparse.Namespace) -> int:
    folder = Path(args.folder)
    if not folder.is_dir():
        print(f"ERROR: folder does not exist: {folder}", file=sys.stderr)
        return 2
    if fitz is None:
        print("ERROR: PyMuPDF is missing. Run Repair_or_Update_Environment.bat.", file=sys.stderr)
        return 2

    options = ScanOptions(
        use_ocr=args.ocr,
        ocr_dpi=args.dpi,
        tesseract_path=args.tesseract or "",
        stop_after_strong_evidence=not args.scan_all_pages,
    )
    if options.use_ocr:
        try:
            configure_tesseract(options.tesseract_path)
        except Exception as exc:
            print(f"ERROR: {exc}", file=sys.stderr)
            return 2

    pdfs = discover_pdfs(folder, not args.no_recursive)
    if not pdfs:
        print("No PDF files found.")
        return 0

    results: list[ScanResult] = []
    for index, pdf_path in enumerate(pdfs, start=1):
        print(f"[{index}/{len(pdfs)}] {pdf_path.name}")

        def page_callback(page: int, total_pages: int, _name: str) -> None:
            print(f"    page {page}/{total_pages}", end="\r", flush=True)

        result = scan_pdf(pdf_path, options, page_callback=page_callback)
        print(f"    {result.status} | score {result.score} | pages {result.matched_pages or '-'}")
        results.append(result)

    output = Path(args.output) if args.output else folder / "fracture_data_scan_results.xlsx"
    settings = {
        "Source Folder": str(folder),
        "Recursive": not args.no_recursive,
        "OCR Enabled": args.ocr,
        "OCR DPI": args.dpi,
        "Stop After Strong Evidence": not args.scan_all_pages,
    }
    export_results(results, output, settings)
    print(f"Saved: {output}")
    return 0


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Scan PDF well files for fracture-test and fracture-treatment data.")
    parser.add_argument("--folder", help="Run in command-line mode on this folder. Omit to launch the GUI.")
    parser.add_argument("--output", help="Output .xlsx or .csv path for command-line mode.")
    parser.add_argument("--ocr", action="store_true", help="OCR image-only pages in command-line mode.")
    parser.add_argument("--dpi", type=int, default=DEFAULT_OCR_DPI, help="OCR render DPI (default: 150).")
    parser.add_argument("--tesseract", default="", help="Path to tesseract.exe if it is not on PATH.")
    parser.add_argument("--workers", type=int, default=1, help="Reserved for future CLI parallel scans; GUI supports 1-4 workers.")
    parser.add_argument("--no-recursive", action="store_true", help="Do not scan subfolders.")
    parser.add_argument("--scan-all-pages", action="store_true", help="Do not stop after strong evidence is found.")
    return parser


def main() -> int:
    relaunched = relaunch_through_managed_environment()
    if relaunched is not None:
        return relaunched

    args = build_arg_parser().parse_args()
    if args.folder:
        return run_cli(args)

    missing: list[str] = []
    if fitz is None:
        missing.append("PyMuPDF")
    if Workbook is None:
        missing.append("openpyxl")
    if missing:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror(
            APP_TITLE,
            "Required package(s) are missing: "
            + ", ".join(missing)
            + "\n\nRun Repair_or_Update_Environment.bat, then start the scanner again.",
        )
        root.destroy()
        return 2

    root = tk.Tk()
    app = FractureScannerApp(root)
    root.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
