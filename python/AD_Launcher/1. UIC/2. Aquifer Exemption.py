"""
ODW Well Monthly Production & Injection Viewer
================================================

A Tkinter desktop application for running a parameterized Oracle query against
ODW. Users can enter one or more engineering strings and an inclusive date
range, review paged results, search/sort the data, and export the current view.

Python 3.10+ is recommended.
"""

from __future__ import annotations

import csv
import json
import logging
from logging.handlers import RotatingFileHandler
import math
import os
from pathlib import Path
import queue
import re
import sys
import threading
import time
from dataclasses import dataclass
from datetime import date, datetime, time as datetime_time, timedelta
from decimal import Decimal
from typing import Any, Callable, Iterable, Sequence
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import oracledb
except ImportError:  # Keep the GUI usable long enough to show install guidance.
    oracledb = None  # type: ignore[assignment]

try:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None  # type: ignore[assignment,misc]
    WriteOnlyCell = None  # type: ignore[assignment,misc]
    Font = None  # type: ignore[assignment,misc]
    get_column_letter = None  # type: ignore[assignment,misc]


APP_NAME = "ODW Well Monthly Production & Injection Viewer"
APP_VERSION = "1.0.0"
DATE_FORMAT = "%Y-%m-%d"
FETCH_BATCH_SIZE = 2_000
MAX_ALLOWED_ROWS = 1_000_000
DEFAULT_MAX_ROWS = 100_000
DEFAULT_PAGE_SIZE = 1_000
PAGE_SIZE_OPTIONS = (250, 500, 1_000, 2_500, 5_000)
DEFAULT_ENGINEERING_STRINGS = "AE-LHH, LHLW, LH-DPSC"
DEFAULT_START_DATE = "2020-05-01"
DEFAULT_END_DATE = "2025-04-30"
DEFAULT_DSN = os.getenv("ODW_DB_DSN", "ODW")
DEFAULT_USERNAME = os.getenv("ODW_DB_USER", "")
DEFAULT_PASSWORD = os.getenv("ODW_DB_PASSWORD", "")

APP_DATA_DIR = Path(os.getenv("LOCALAPPDATA", str(Path.home()))) / "ODWWellMonthlyViewer"
SETTINGS_PATH = APP_DATA_DIR / "settings.json"
LOG_PATH = APP_DATA_DIR / "odw_well_monthly_viewer.log"


SQL_TEMPLATE = """
WITH st AS (
    -- Return exactly one latest active completion state per well.
    SELECT well_fac_id,
           cmpl_state_type_desc AS well_state
    FROM (
        SELECT well_fac_id,
               cmpl_fac_id,
               cmpl_nme,
               cmpl_state_type_desc,
               ROW_NUMBER() OVER (
                   PARTITION BY well_fac_id
                   ORDER BY cmpl_fac_id DESC, cmpl_nme DESC
               ) AS rn
        FROM cmpl_dmn
        WHERE actv_indc = 'Y'
          AND cmpl_state_type_cde NOT IN ('PRPO', 'FUTR')
          AND engr_strg_nme IN ({engineering_binds})
    )
    WHERE rn = 1
)
SELECT wd.well_nme                         AS "WELL NAME",
       wd.well_api_nbr                     AS "WELL API",
       cd.engr_strg_nme                    AS "ENGINEERING STRING",
       cf.eftv_dttm                        AS "DATE",
       cf.aloc_oil_prod_dly_rte_qty        AS "OIL PROD BBL",
       cf.aloc_wtr_prod_dly_rte_qty        AS "WATER PROD BBL",
       cf.aloc_gas_prod_dly_rte_qty        AS "GAS PROD MCF",
       cf.aloc_stm_inj_dly_rte_qty         AS "STEAM INJ BBL",
       cf.aloc_wtr_inj_on_days_rte_qty     AS "WATER INJ BBL",
       st.well_state                       AS "STATUS"
FROM well_dmn wd
JOIN cmpl_dmn cd
  ON wd.well_fac_id = cd.well_fac_id
JOIN st
  ON wd.well_fac_id = st.well_fac_id
JOIN cmpl_mnly_fact cf
  ON cd.cmpl_fac_id = cf.cmpl_fac_id
WHERE cd.actv_indc = 'Y'
  AND wd.actv_indc = 'Y'
  AND cd.engr_strg_nme IN ({engineering_binds})
  AND cd.cmpl_state_type_cde NOT IN ('PRPO', 'FUTR')
  AND cf.eftv_dttm >= :start_date
  AND cf.eftv_dttm < :end_date_exclusive
ORDER BY wd.well_api_nbr,
         cf.eftv_dttm
""".strip()


@dataclass(frozen=True)
class ConnectionSettings:
    username: str
    password: str
    dsn: str


@dataclass(frozen=True)
class QueryRequest:
    connection: ConnectionSettings
    engineering_strings: tuple[str, ...]
    start_date: date
    end_date: date
    max_rows: int

    @property
    def end_date_exclusive(self) -> datetime:
        return datetime.combine(self.end_date + timedelta(days=1), datetime_time.min)


@dataclass(frozen=True)
class QueryResult:
    columns: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]
    truncated: bool
    elapsed_seconds: float


@dataclass(frozen=True)
class ExportResult:
    output_path: Path
    row_count: int
    elapsed_seconds: float


def configure_logging() -> logging.Logger:
    APP_DATA_DIR.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("odw_well_monthly_viewer")
    logger.setLevel(logging.INFO)
    if not logger.handlers:
        handler = RotatingFileHandler(
            LOG_PATH,
            maxBytes=1_500_000,
            backupCount=3,
            encoding="utf-8",
        )
        handler.setFormatter(
            logging.Formatter("%(asctime)s | %(levelname)s | %(threadName)s | %(message)s")
        )
        logger.addHandler(handler)
    return logger


LOGGER = configure_logging()


def initialize_oracle_driver() -> str:
    """Initialize python-oracledb in auto/thin/thick mode and return a label."""
    if oracledb is None:
        return "Not installed"

    requested_mode = os.getenv("ODW_ORACLE_MODE", "auto").strip().lower()
    lib_dir = os.getenv("ODW_ORACLE_CLIENT_LIB_DIR", "").strip()

    if requested_mode not in {"auto", "thin", "thick"}:
        LOGGER.warning("Unknown ODW_ORACLE_MODE=%s; using auto mode", requested_mode)
        requested_mode = "auto"

    if requested_mode in {"auto", "thick"}:
        try:
            if lib_dir:
                oracledb.init_oracle_client(lib_dir=lib_dir)
            else:
                oracledb.init_oracle_client()
        except Exception as exc:
            if requested_mode == "thick":
                LOGGER.exception("Oracle Thick mode initialization failed")
                return f"Thick mode error: {friendly_exception(exc)}"
            LOGGER.info("Oracle Client not initialized; continuing in Thin mode: %s", exc)

    try:
        return "Thin" if oracledb.is_thin_mode() else "Thick"
    except Exception:
        return "Available"


def load_settings() -> dict[str, Any]:
    try:
        if SETTINGS_PATH.exists():
            data = json.loads(SETTINGS_PATH.read_text(encoding="utf-8"))
            return data if isinstance(data, dict) else {}
    except Exception:
        LOGGER.exception("Could not load settings")
    return {}


def save_settings(data: dict[str, Any]) -> None:
    try:
        APP_DATA_DIR.mkdir(parents=True, exist_ok=True)
        temp_path = SETTINGS_PATH.with_suffix(".tmp")
        temp_path.write_text(json.dumps(data, indent=2), encoding="utf-8")
        temp_path.replace(SETTINGS_PATH)
    except Exception:
        LOGGER.exception("Could not save settings")


def parse_engineering_strings(raw_value: str) -> tuple[str, ...]:
    """Parse comma, semicolon, or newline separated engineering strings."""
    candidates = re.split(r"[,;\n]+", raw_value)
    values: list[str] = []
    seen: set[str] = set()

    for candidate in candidates:
        value = candidate.strip().strip("'\"").strip()
        if not value:
            continue
        if len(value) > 100:
            raise ValueError(f"Engineering string is too long: {value[:40]}...")
        key = value.casefold()
        if key not in seen:
            seen.add(key)
            values.append(value)

    if not values:
        raise ValueError("Enter at least one engineering string.")
    if len(values) > 100:
        raise ValueError("No more than 100 engineering strings can be queried at once.")
    return tuple(values)


def parse_iso_date(raw_value: str, field_name: str) -> date:
    try:
        return datetime.strptime(raw_value.strip(), DATE_FORMAT).date()
    except ValueError as exc:
        raise ValueError(f"{field_name} must use YYYY-MM-DD format.") from exc


def parse_max_rows(raw_value: str) -> int:
    try:
        value = int(raw_value.replace(",", "").strip())
    except ValueError as exc:
        raise ValueError("Max rows must be a whole number.") from exc
    if value < 1 or value > MAX_ALLOWED_ROWS:
        raise ValueError(f"Max rows must be between 1 and {MAX_ALLOWED_ROWS:,}.")
    return value


def build_query(
    engineering_strings: Sequence[str],
    start_date: date,
    end_date: date,
) -> tuple[str, dict[str, Any]]:
    if not engineering_strings:
        raise ValueError("At least one engineering string is required.")
    if end_date < start_date:
        raise ValueError("End date cannot be earlier than start date.")

    bind_names = [f"eng_{index}" for index in range(len(engineering_strings))]
    bind_sql = ", ".join(f":{name}" for name in bind_names)
    sql = SQL_TEMPLATE.format(engineering_binds=bind_sql)

    params: dict[str, Any] = {
        name: value for name, value in zip(bind_names, engineering_strings)
    }
    params["start_date"] = datetime.combine(start_date, datetime_time.min)
    params["end_date_exclusive"] = datetime.combine(
        end_date + timedelta(days=1), datetime_time.min
    )
    return sql, params


def friendly_exception(exc: BaseException) -> str:
    if oracledb is not None and isinstance(exc, oracledb.Error):
        detail = exc.args[0] if exc.args else exc
        code = getattr(detail, "code", None)
        message = str(getattr(detail, "message", detail)).strip()
        return f"Oracle error {code}: {message}" if code else message
    message = str(exc).strip()
    return f"{type(exc).__name__}: {message}" if message else type(exc).__name__


def open_connection(settings: ConnectionSettings):
    if oracledb is None:
        raise RuntimeError(
            "python-oracledb is not installed. Run: python -m pip install -r requirements.txt"
        )
    return oracledb.connect(
        user=settings.username,
        password=settings.password,
        dsn=settings.dsn,
    )


def test_database_connection(
    settings: ConnectionSettings,
    progress: Callable[[str], None],
) -> str:
    progress("Connecting to Oracle...")
    connection = open_connection(settings)
    cursor = None
    try:
        cursor = connection.cursor()
        cursor.execute("SELECT 1 FROM dual")
        cursor.fetchone()
        version = getattr(connection, "version", "unknown")
        return str(version)
    finally:
        if cursor is not None:
            cursor.close()
        connection.close()


def execute_monthly_query(
    request: QueryRequest,
    progress: Callable[[str], None],
) -> QueryResult:
    sql, params = build_query(
        request.engineering_strings,
        request.start_date,
        request.end_date,
    )
    started = time.perf_counter()
    LOGGER.info(
        "Running query | user=%s | dsn=%s | engineering_strings=%s | start=%s | end=%s | max_rows=%s",
        request.connection.username,
        request.connection.dsn,
        request.engineering_strings,
        request.start_date,
        request.end_date,
        request.max_rows,
    )

    progress("Connecting to Oracle...")
    connection = open_connection(request.connection)
    cursor = None
    rows: list[tuple[Any, ...]] = []
    truncated = False

    try:
        cursor = connection.cursor()
        cursor.arraysize = FETCH_BATCH_SIZE
        try:
            cursor.prefetchrows = FETCH_BATCH_SIZE
        except Exception:
            pass

        progress("Executing query...")
        cursor.execute(sql, params)
        columns = tuple(str(item[0]) for item in (cursor.description or ()))

        next_progress_count = 10_000
        while True:
            remaining = request.max_rows + 1 - len(rows)
            if remaining <= 0:
                truncated = True
                break

            batch = cursor.fetchmany(min(FETCH_BATCH_SIZE, remaining))
            if not batch:
                break
            rows.extend(tuple(row) for row in batch)

            if len(rows) >= next_progress_count:
                progress(f"Fetched {len(rows):,} rows...")
                next_progress_count += 10_000

            if len(rows) > request.max_rows:
                truncated = True
                rows = rows[: request.max_rows]
                break

        elapsed = time.perf_counter() - started
        LOGGER.info(
            "Query complete | rows=%s | truncated=%s | elapsed=%.2fs",
            len(rows),
            truncated,
            elapsed,
        )
        return QueryResult(columns, tuple(rows), truncated, elapsed)
    finally:
        if cursor is not None:
            cursor.close()
        connection.close()


def normalize_export_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return value
    if isinstance(value, Decimal):
        return float(value)
    return value


def export_rows(
    output_path: Path,
    columns: Sequence[str],
    rows: Sequence[Sequence[Any]],
    progress: Callable[[str], None],
) -> ExportResult:
    started = time.perf_counter()
    suffix = output_path.suffix.lower()

    if suffix == ".csv":
        progress("Writing CSV...")
        with output_path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(columns)
            for index, row in enumerate(rows, start=1):
                writer.writerow(normalize_export_value(value) for value in row)
                if index % 10_000 == 0:
                    progress(f"Exported {index:,} rows...")
    elif suffix == ".xlsx":
        if Workbook is None or WriteOnlyCell is None or Font is None or get_column_letter is None:
            raise RuntimeError(
                "Excel export requires openpyxl. Run: python -m pip install openpyxl"
            )

        progress("Writing Excel workbook...")
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet("Monthly Data")
        worksheet.freeze_panes = "A2"

        header_cells = []
        for column in columns:
            cell = WriteOnlyCell(worksheet, value=column)
            cell.font = Font(bold=True)
            header_cells.append(cell)
        worksheet.append(header_cells)

        sample_widths = [len(str(column)) for column in columns]
        for index, row in enumerate(rows, start=1):
            excel_row = []
            for column_index, value in enumerate(row):
                normalized = normalize_export_value(value)
                cell = WriteOnlyCell(worksheet, value=normalized)
                if isinstance(normalized, (datetime, date)):
                    cell.number_format = "yyyy-mm-dd"
                excel_row.append(cell)
                if index <= 2_000:
                    sample_widths[column_index] = min(
                        45,
                        max(sample_widths[column_index], len(str(normalized or ""))),
                    )
            worksheet.append(excel_row)
            if index % 10_000 == 0:
                progress(f"Exported {index:,} rows...")

        for column_index, width in enumerate(sample_widths, start=1):
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                45, max(10, width + 2)
            )

        if rows:
            worksheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
            )
        workbook.save(output_path)
    else:
        raise ValueError("Choose a .csv or .xlsx output file.")

    elapsed = time.perf_counter() - started
    LOGGER.info(
        "Export complete | path=%s | rows=%s | elapsed=%.2fs",
        output_path,
        len(rows),
        elapsed,
    )
    return ExportResult(output_path, len(rows), elapsed)


def display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime(DATE_FORMAT)
    if isinstance(value, date):
        return value.strftime(DATE_FORMAT)
    if isinstance(value, Decimal):
        value = float(value)
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return f"{value:,.0f}"
        return f"{value:,.2f}".rstrip("0").rstrip(".")
    if isinstance(value, int):
        return f"{value:,}"
    return str(value)


def sortable_value(value: Any) -> tuple[int, Any]:
    if isinstance(value, datetime):
        return (0, value.timestamp())
    if isinstance(value, date):
        return (0, value.toordinal())
    if isinstance(value, Decimal):
        return (1, float(value))
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return (1, float(value))
    return (2, str(value).casefold())


class ODWWellMonthlyViewer(tk.Tk):
    COLORS = {
        "navy": "#17365D",
        "blue": "#2F75B5",
        "light_blue": "#D9EAF7",
        "panel": "#F7F9FC",
        "border": "#C9D4E2",
        "text": "#1F2937",
        "muted": "#5B6573",
        "white": "#FFFFFF",
        "odd": "#F5F8FC",
        "warning": "#9A6700",
    }

    def __init__(self, driver_mode: str) -> None:
        super().__init__()
        self.driver_mode = driver_mode
        self.settings_data = load_settings()
        self.event_queue: queue.Queue[tuple[str, str, Any]] = queue.Queue()
        self.worker_thread: threading.Thread | None = None
        self.busy_operation: str | None = None
        self.columns: tuple[str, ...] = ()
        self.all_rows: list[tuple[Any, ...]] = []
        self.filtered_rows: list[tuple[Any, ...]] = []
        self.current_page = 0
        self.sort_column_index: int | None = None
        self.sort_reverse = False
        self.search_after_id: str | None = None
        self.last_export_dir = str(
            self.settings_data.get("last_export_dir", Path.home() / "Documents")
        )

        self.title(f"{APP_NAME} v{APP_VERSION}")
        self.minsize(1_050, 680)
        geometry = str(self.settings_data.get("geometry", "1380x820"))
        if re.fullmatch(r"\d+x\d+(?:[+-]\d+){0,2}", geometry):
            self.geometry(geometry)
        else:
            self.geometry("1380x820")

        self._configure_style()
        self._create_variables()
        self._build_ui()
        self._bind_shortcuts()
        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self.after(100, self._poll_worker_events)

    def _configure_style(self) -> None:
        style = ttk.Style(self)
        available = style.theme_names()
        if "vista" in available and sys.platform.startswith("win"):
            style.theme_use("vista")
        elif "clam" in available:
            style.theme_use("clam")

        style.configure("TFrame", background=self.COLORS["panel"])
        style.configure("TLabel", background=self.COLORS["panel"], foreground=self.COLORS["text"])
        style.configure("TLabelframe", background=self.COLORS["panel"])
        style.configure(
            "TLabelframe.Label",
            background=self.COLORS["panel"],
            foreground=self.COLORS["navy"],
            font=("Segoe UI", 9, "bold"),
        )
        style.configure("TButton", font=("Segoe UI", 9), padding=(10, 5))
        style.configure(
            "Accent.TButton",
            font=("Segoe UI", 9, "bold"),
            foreground=self.COLORS["white"],
            background=self.COLORS["blue"],
            padding=(14, 6),
        )
        style.map(
            "Accent.TButton",
            background=[("active", "#245F93"), ("disabled", "#AAB7C4")],
            foreground=[("disabled", "#EFF3F6")],
        )
        style.configure(
            "Treeview",
            font=("Segoe UI", 9),
            rowheight=24,
            background=self.COLORS["white"],
            fieldbackground=self.COLORS["white"],
            bordercolor=self.COLORS["border"],
        )
        style.configure(
            "Treeview.Heading",
            font=("Segoe UI", 9, "bold"),
            background="#E8EEF5",
            foreground=self.COLORS["navy"],
            padding=(6, 6),
        )
        style.map("Treeview", background=[("selected", self.COLORS["light_blue"])])
        style.configure("Status.TLabel", font=("Segoe UI", 9), foreground=self.COLORS["muted"])
        style.configure("Summary.TLabel", font=("Segoe UI", 9, "bold"), foreground=self.COLORS["navy"])

    def _create_variables(self) -> None:
        self.username_var = tk.StringVar(
            value=str(self.settings_data.get("username", DEFAULT_USERNAME))
        )
        self.password_var = tk.StringVar(value=DEFAULT_PASSWORD)
        self.dsn_var = tk.StringVar(value=str(self.settings_data.get("dsn", DEFAULT_DSN)))
        self.show_password_var = tk.BooleanVar(value=False)
        self.engineering_var = tk.StringVar(
            value=str(
                self.settings_data.get(
                    "engineering_strings", DEFAULT_ENGINEERING_STRINGS
                )
            )
        )
        self.start_date_var = tk.StringVar(
            value=str(self.settings_data.get("start_date", DEFAULT_START_DATE))
        )
        self.end_date_var = tk.StringVar(
            value=str(self.settings_data.get("end_date", DEFAULT_END_DATE))
        )
        self.max_rows_var = tk.StringVar(
            value=str(self.settings_data.get("max_rows", DEFAULT_MAX_ROWS))
        )
        self.page_size_var = tk.StringVar(
            value=str(self.settings_data.get("page_size", DEFAULT_PAGE_SIZE))
        )
        self.search_var = tk.StringVar()
        self.status_var = tk.StringVar(value="Ready.")
        self.summary_var = tk.StringVar(value="No query results loaded.")
        self.page_var = tk.StringVar(value="Page 0 of 0")

    def _build_ui(self) -> None:
        self.configure(background=self.COLORS["panel"])

        header = tk.Frame(self, background=self.COLORS["navy"], height=72)
        header.pack(fill="x")
        header.pack_propagate(False)

        title_frame = tk.Frame(header, background=self.COLORS["navy"])
        title_frame.pack(side="left", fill="y", padx=18)
        tk.Label(
            title_frame,
            text="ODW WELL MONTHLY DATA",
            font=("Segoe UI", 16, "bold"),
            foreground=self.COLORS["white"],
            background=self.COLORS["navy"],
        ).pack(anchor="w", pady=(10, 0))
        tk.Label(
            title_frame,
            text="Production and injection rates by engineering string and date range",
            font=("Segoe UI", 9),
            foreground="#D9E5F2",
            background=self.COLORS["navy"],
        ).pack(anchor="w")

        driver_label = tk.Label(
            header,
            text=f"Oracle mode: {self.driver_mode}",
            font=("Segoe UI", 9),
            foreground="#D9E5F2",
            background=self.COLORS["navy"],
        )
        driver_label.pack(side="right", padx=18)

        main = ttk.Frame(self, padding=(12, 10, 12, 8))
        main.pack(fill="both", expand=True)
        main.columnconfigure(0, weight=1)
        main.rowconfigure(2, weight=1)

        self._build_inputs(main)
        self._build_result_toolbar(main)
        self._build_results(main)
        self._build_status_bar(main)

    def _build_inputs(self, parent: ttk.Frame) -> None:
        input_frame = ttk.LabelFrame(parent, text="Connection and Query Inputs", padding=10)
        input_frame.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        input_frame.columnconfigure(0, weight=1)

        connection_row = ttk.Frame(input_frame)
        connection_row.grid(row=0, column=0, sticky="ew")
        connection_row.columnconfigure(7, weight=1)

        ttk.Label(connection_row, text="Username").grid(
            row=0, column=0, sticky="w", padx=(0, 5)
        )
        self.username_entry = ttk.Entry(
            connection_row, textvariable=self.username_var, width=15
        )
        self.username_entry.grid(row=0, column=1, sticky="ew", padx=(0, 10))

        ttk.Label(connection_row, text="Password").grid(
            row=0, column=2, sticky="w", padx=(0, 5)
        )
        self.password_entry = ttk.Entry(
            connection_row,
            textvariable=self.password_var,
            show="*",
            width=15,
        )
        self.password_entry.grid(row=0, column=3, sticky="ew", padx=(0, 3))
        ttk.Checkbutton(
            connection_row,
            text="Show",
            variable=self.show_password_var,
            command=self._toggle_password,
        ).grid(row=0, column=4, sticky="w", padx=(0, 10))

        ttk.Label(connection_row, text="DSN").grid(
            row=0, column=5, sticky="w", padx=(0, 5)
        )
        self.dsn_entry = ttk.Entry(connection_row, textvariable=self.dsn_var, width=17)
        self.dsn_entry.grid(row=0, column=6, sticky="ew", padx=(0, 10))

        ttk.Frame(connection_row).grid(row=0, column=7, sticky="ew")

        self.test_button = ttk.Button(
            connection_row,
            text="Test Connection",
            command=self._test_connection,
        )
        self.test_button.grid(row=0, column=8, sticky="e", padx=(0, 6))

        self.sql_button = ttk.Button(connection_row, text="View SQL", command=self._show_sql)
        self.sql_button.grid(row=0, column=9, sticky="e")

        ttk.Separator(input_frame, orient="horizontal").grid(
            row=1, column=0, sticky="ew", pady=9
        )

        engineering_row = ttk.Frame(input_frame)
        engineering_row.grid(row=2, column=0, sticky="ew")
        engineering_row.columnconfigure(1, weight=1)
        ttk.Label(engineering_row, text="Engineering string(s)").grid(
            row=0, column=0, sticky="w", padx=(0, 5)
        )
        self.engineering_entry = ttk.Entry(
            engineering_row,
            textvariable=self.engineering_var,
        )
        self.engineering_entry.grid(row=0, column=1, sticky="ew")

        date_row = ttk.Frame(input_frame)
        date_row.grid(row=3, column=0, sticky="ew", pady=(8, 0))
        date_row.columnconfigure(9, weight=1)

        ttk.Label(date_row, text="Start date").grid(
            row=0, column=0, sticky="w", padx=(0, 5)
        )
        self.start_date_entry = ttk.Entry(
            date_row,
            textvariable=self.start_date_var,
            width=12,
        )
        self.start_date_entry.grid(row=0, column=1, sticky="w", padx=(0, 14))

        ttk.Label(date_row, text="End date (inclusive)").grid(
            row=0, column=2, sticky="w", padx=(0, 5)
        )
        self.end_date_entry = ttk.Entry(
            date_row,
            textvariable=self.end_date_var,
            width=12,
        )
        self.end_date_entry.grid(row=0, column=3, sticky="w", padx=(0, 14))

        ttk.Label(date_row, text="Max rows").grid(
            row=0, column=4, sticky="w", padx=(0, 5)
        )
        self.max_rows_entry = ttk.Entry(
            date_row,
            textvariable=self.max_rows_var,
            width=10,
        )
        self.max_rows_entry.grid(row=0, column=5, sticky="w", padx=(0, 8))

        ttk.Label(
            date_row,
            text="Dates: YYYY-MM-DD",
            foreground=self.COLORS["muted"],
            font=("Segoe UI", 8),
        ).grid(row=0, column=6, sticky="w", padx=(0, 8))

        ttk.Frame(date_row).grid(row=0, column=9, sticky="ew")
        self.run_button = ttk.Button(
            date_row,
            text="Run Query",
            style="Accent.TButton",
            command=self._run_query,
        )
        self.run_button.grid(row=0, column=10, sticky="e")

        help_text = (
            "Separate multiple engineering strings with commas, semicolons, or new lines. "
            "The selected end date is included."
        )
        ttk.Label(
            input_frame,
            text=help_text,
            foreground=self.COLORS["muted"],
            font=("Segoe UI", 8),
        ).grid(row=4, column=0, sticky="w", pady=(6, 0))

    def _build_result_toolbar(self, parent: ttk.Frame) -> None:
        toolbar = ttk.Frame(parent)
        toolbar.grid(row=1, column=0, sticky="ew", pady=(0, 6))
        toolbar.columnconfigure(1, weight=1)

        ttk.Label(toolbar, text="Search results").grid(row=0, column=0, sticky="w", padx=(0, 5))
        self.search_entry = ttk.Entry(toolbar, textvariable=self.search_var)
        self.search_entry.grid(row=0, column=1, sticky="ew", padx=(0, 5))
        self.search_var.trace_add("write", self._schedule_filter)

        ttk.Button(toolbar, text="Clear", command=lambda: self.search_var.set("")).grid(
            row=0, column=2, padx=(0, 12)
        )

        ttk.Label(toolbar, text="Rows per page").grid(row=0, column=3, padx=(0, 5))
        self.page_size_combo = ttk.Combobox(
            toolbar,
            textvariable=self.page_size_var,
            values=[str(value) for value in PAGE_SIZE_OPTIONS],
            width=7,
            state="readonly",
        )
        self.page_size_combo.grid(row=0, column=4, padx=(0, 12))
        self.page_size_combo.bind("<<ComboboxSelected>>", lambda _event: self._page_size_changed())

        self.copy_button = ttk.Button(toolbar, text="Copy Selected", command=self._copy_selected)
        self.copy_button.grid(row=0, column=5, padx=(0, 5))
        self.copy_button.state(["disabled"])

        self.export_button = ttk.Button(toolbar, text="Export View", command=self._export_view)
        self.export_button.grid(row=0, column=6)
        self.export_button.state(["disabled"])

    def _build_results(self, parent: ttk.Frame) -> None:
        result_frame = ttk.Frame(parent)
        result_frame.grid(row=2, column=0, sticky="nsew")
        result_frame.columnconfigure(0, weight=1)
        result_frame.rowconfigure(1, weight=1)

        summary_frame = ttk.Frame(result_frame)
        summary_frame.grid(row=0, column=0, sticky="ew", pady=(0, 4))
        summary_frame.columnconfigure(0, weight=1)
        ttk.Label(summary_frame, textvariable=self.summary_var, style="Summary.TLabel").grid(
            row=0, column=0, sticky="w"
        )

        table_frame = ttk.Frame(result_frame)
        table_frame.grid(row=1, column=0, sticky="nsew")
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)

        self.tree = ttk.Treeview(table_frame, show="headings", selectmode="extended")
        vertical_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        horizontal_scroll = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(
            yscrollcommand=vertical_scroll.set,
            xscrollcommand=horizontal_scroll.set,
        )
        self.tree.grid(row=0, column=0, sticky="nsew")
        vertical_scroll.grid(row=0, column=1, sticky="ns")
        horizontal_scroll.grid(row=1, column=0, sticky="ew")
        self.tree.tag_configure("even", background=self.COLORS["white"])
        self.tree.tag_configure("odd", background=self.COLORS["odd"])
        self.tree.bind("<<TreeviewSelect>>", self._tree_selection_changed)
        self.tree.bind("<Control-c>", lambda _event: self._copy_selected())
        self.tree.bind("<Button-3>", self._show_tree_context_menu)

        pager = ttk.Frame(result_frame)
        pager.grid(row=2, column=0, sticky="ew", pady=(6, 0))
        pager.columnconfigure(4, weight=1)

        self.first_button = ttk.Button(pager, text="|<", width=4, command=self._first_page)
        self.previous_button = ttk.Button(pager, text="<", width=4, command=self._previous_page)
        self.next_button = ttk.Button(pager, text=">", width=4, command=self._next_page)
        self.last_button = ttk.Button(pager, text=">|", width=4, command=self._last_page)
        self.first_button.grid(row=0, column=0, padx=(0, 3))
        self.previous_button.grid(row=0, column=1, padx=(0, 8))
        ttk.Label(pager, textvariable=self.page_var).grid(row=0, column=2, padx=(0, 8))
        self.next_button.grid(row=0, column=3, padx=(0, 3))
        self.last_button.grid(row=0, column=4, sticky="w")

        self._update_pager_state()

    def _build_status_bar(self, parent: ttk.Frame) -> None:
        status_frame = ttk.Frame(parent)
        status_frame.grid(row=3, column=0, sticky="ew", pady=(7, 0))
        status_frame.columnconfigure(0, weight=1)
        ttk.Label(status_frame, textvariable=self.status_var, style="Status.TLabel").grid(
            row=0, column=0, sticky="w"
        )
        self.progress = ttk.Progressbar(status_frame, mode="indeterminate", length=180)
        self.progress.grid(row=0, column=1, sticky="e")

    def _bind_shortcuts(self) -> None:
        self.bind("<F5>", lambda _event: self._run_query())
        self.bind("<Control-Return>", lambda _event: self._run_query())
        self.bind("<Control-e>", lambda _event: self._export_view())
        self.bind("<Control-f>", lambda _event: self.search_entry.focus_set())
        self.bind("<Escape>", lambda _event: self.search_var.set(""))

    def _toggle_password(self) -> None:
        self.password_entry.configure(show="" if self.show_password_var.get() else "*")

    def _collect_connection_settings(self) -> ConnectionSettings:
        username = self.username_var.get().strip()
        password = self.password_var.get()
        dsn = self.dsn_var.get().strip()
        if not username:
            raise ValueError("Enter the Oracle username.")
        if not password:
            raise ValueError("Enter the Oracle password.")
        if not dsn:
            raise ValueError("Enter the DSN or TNS alias.")
        return ConnectionSettings(username=username, password=password, dsn=dsn)

    def _collect_query_request(self) -> QueryRequest:
        connection = self._collect_connection_settings()
        engineering_strings = parse_engineering_strings(self.engineering_var.get())
        start_date = parse_iso_date(self.start_date_var.get(), "Start date")
        end_date = parse_iso_date(self.end_date_var.get(), "End date")
        if end_date < start_date:
            raise ValueError("End date cannot be earlier than start date.")
        max_rows = parse_max_rows(self.max_rows_var.get())
        return QueryRequest(
            connection=connection,
            engineering_strings=engineering_strings,
            start_date=start_date,
            end_date=end_date,
            max_rows=max_rows,
        )

    def _start_worker(
        self,
        operation: str,
        target: Callable[..., Any],
        *args: Any,
    ) -> None:
        if self.busy_operation is not None:
            messagebox.showinfo(
                "Operation in progress",
                f"Please wait for {self.busy_operation} to finish.",
                parent=self,
            )
            return

        self.busy_operation = operation
        self.run_button.state(["disabled"])
        self.test_button.state(["disabled"])
        self.export_button.state(["disabled"])
        self.progress.start(12)
        self.configure(cursor="watch")

        def progress(message: str) -> None:
            self.event_queue.put(("progress", operation, message))

        def runner() -> None:
            try:
                result = target(*args, progress)
            except Exception as exc:
                LOGGER.exception("Background operation failed: %s", operation)
                self.event_queue.put(("error", operation, friendly_exception(exc)))
            else:
                self.event_queue.put(("success", operation, result))

        self.worker_thread = threading.Thread(
            target=runner,
            name=f"{operation}-worker",
            daemon=True,
        )
        self.worker_thread.start()

    def _finish_worker(self) -> None:
        self.busy_operation = None
        self.worker_thread = None
        self.progress.stop()
        self.configure(cursor="")
        self.run_button.state(["!disabled"])
        self.test_button.state(["!disabled"])
        if self.filtered_rows:
            self.export_button.state(["!disabled"])

    def _poll_worker_events(self) -> None:
        try:
            while True:
                event_type, operation, payload = self.event_queue.get_nowait()
                if event_type == "progress":
                    if operation == self.busy_operation:
                        self.status_var.set(str(payload))
                elif event_type == "error":
                    self._finish_worker()
                    self.status_var.set(f"{operation} failed.")
                    messagebox.showerror(
                        f"{operation.title()} Failed",
                        str(payload),
                        parent=self,
                    )
                elif event_type == "success":
                    self._finish_worker()
                    if operation == "connection test":
                        self.status_var.set("Connection test succeeded.")
                        messagebox.showinfo(
                            "Connection Successful",
                            f"Connected to Oracle Database {payload}.",
                            parent=self,
                        )
                    elif operation == "query":
                        self._accept_query_result(payload)
                    elif operation == "export":
                        self._accept_export_result(payload)
        except queue.Empty:
            pass
        finally:
            try:
                self.after(100, self._poll_worker_events)
            except tk.TclError:
                pass

    def _test_connection(self) -> None:
        try:
            settings = self._collect_connection_settings()
        except ValueError as exc:
            messagebox.showwarning("Check Inputs", str(exc), parent=self)
            return
        self.status_var.set("Testing connection...")
        self._start_worker("connection test", test_database_connection, settings)

    def _run_query(self) -> None:
        try:
            request = self._collect_query_request()
        except ValueError as exc:
            messagebox.showwarning("Check Inputs", str(exc), parent=self)
            return
        self.status_var.set("Starting query...")
        self._start_worker("query", execute_monthly_query, request)

    def _accept_query_result(self, result: QueryResult) -> None:
        self.columns = result.columns
        self.all_rows = list(result.rows)
        self.filtered_rows = list(result.rows)
        self.current_page = 0
        self.sort_column_index = None
        self.sort_reverse = False
        self.search_var.set("")
        self._configure_tree_columns()
        self._update_summary(result)
        self._render_page()

        truncation_text = " Results reached the configured row limit." if result.truncated else ""
        self.status_var.set(
            f"Loaded {len(result.rows):,} rows in {result.elapsed_seconds:,.1f} seconds.{truncation_text}"
        )
        if self.filtered_rows:
            self.export_button.state(["!disabled"])
        else:
            self.export_button.state(["disabled"])
        self.copy_button.state(["disabled"])
        self._save_current_settings()

    def _update_summary(self, result: QueryResult) -> None:
        if not result.rows:
            self.summary_var.set("Query completed successfully; no rows matched the inputs.")
            return

        api_index = self._column_index("WELL API")
        date_index = self._column_index("DATE")
        unique_wells = len(
            {row[api_index] for row in result.rows if api_index is not None and row[api_index] is not None}
        ) if api_index is not None else 0

        dates: list[date] = []
        if date_index is not None:
            for row in result.rows:
                value = row[date_index]
                if isinstance(value, datetime):
                    dates.append(value.date())
                elif isinstance(value, date):
                    dates.append(value)

        period = ""
        if dates:
            period = f" | Data period: {min(dates):%Y-%m-%d} to {max(dates):%Y-%m-%d}"
        truncated = " | ROW LIMIT REACHED" if result.truncated else ""
        self.summary_var.set(
            f"Rows: {len(result.rows):,} | Unique wells: {unique_wells:,}{period}{truncated}"
        )

    def _column_index(self, column_name: str) -> int | None:
        try:
            return self.columns.index(column_name)
        except ValueError:
            return None

    def _configure_tree_columns(self) -> None:
        internal_columns = ["__row_number__"] + [
            f"column_{index}" for index in range(len(self.columns))
        ]
        self.tree.configure(columns=internal_columns, displaycolumns=internal_columns)

        self.tree.heading("__row_number__", text="#")
        self.tree.column("__row_number__", width=58, minwidth=50, stretch=False, anchor="e")

        width_map = {
            "WELL NAME": 155,
            "WELL API": 115,
            "ENGINEERING STRING": 145,
            "DATE": 105,
            "OIL PROD BBL": 120,
            "WATER PROD BBL": 135,
            "GAS PROD MCF": 125,
            "STEAM INJ BBL": 125,
            "WATER INJ BBL": 125,
            "STATUS": 135,
        }

        for index, column_name in enumerate(self.columns):
            internal_name = f"column_{index}"
            self.tree.heading(
                internal_name,
                text=column_name,
                command=lambda selected_index=index: self._sort_by_column(selected_index),
            )
            numeric = any(
                token in column_name.upper()
                for token in ("BBL", "MCF", "QTY", "RATE", "VOLUME")
            )
            anchor = "e" if numeric else ("center" if column_name == "DATE" else "w")
            self.tree.column(
                internal_name,
                width=width_map.get(column_name, max(100, len(column_name) * 9)),
                minwidth=70,
                stretch=True,
                anchor=anchor,
            )

    def _schedule_filter(self, *_args: Any) -> None:
        if self.search_after_id is not None:
            try:
                self.after_cancel(self.search_after_id)
            except tk.TclError:
                pass
        self.search_after_id = self.after(300, self._apply_filter)

    def _apply_filter(self) -> None:
        self.search_after_id = None
        term = self.search_var.get().strip().casefold()
        if not term:
            self.filtered_rows = list(self.all_rows)
        else:
            self.status_var.set("Filtering results...")
            self.filtered_rows = [
                row
                for row in self.all_rows
                if term in " | ".join(display_value(value).casefold() for value in row)
            ]

        if self.sort_column_index is not None:
            self._sort_current_rows(
                self.sort_column_index,
                self.sort_reverse,
                update_heading=False,
            )
        self.current_page = 0
        self._render_page()
        if self.all_rows:
            self.summary_var.set(
                f"Showing {len(self.filtered_rows):,} of {len(self.all_rows):,} loaded rows"
            )
        self.status_var.set(f"Filter matched {len(self.filtered_rows):,} rows.")
        if self.filtered_rows:
            self.export_button.state(["!disabled"])
        else:
            self.export_button.state(["disabled"])

    def _sort_by_column(self, column_index: int) -> None:
        reverse = (
            not self.sort_reverse
            if self.sort_column_index == column_index
            else False
        )
        self.sort_column_index = column_index
        self.sort_reverse = reverse
        self._sort_current_rows(column_index, reverse, update_heading=True)
        self.current_page = 0
        self._render_page()

    def _sort_current_rows(
        self,
        column_index: int,
        reverse: bool,
        update_heading: bool,
    ) -> None:
        non_null_rows = [row for row in self.filtered_rows if row[column_index] is not None]
        null_rows = [row for row in self.filtered_rows if row[column_index] is None]
        non_null_rows.sort(
            key=lambda row: sortable_value(row[column_index]),
            reverse=reverse,
        )
        self.filtered_rows = non_null_rows + null_rows

        if update_heading:
            for index, column_name in enumerate(self.columns):
                marker = ""
                if index == column_index:
                    marker = "  v" if reverse else "  ^"
                self.tree.heading(f"column_{index}", text=f"{column_name}{marker}")

    def _page_size(self) -> int:
        try:
            return int(self.page_size_var.get())
        except ValueError:
            return DEFAULT_PAGE_SIZE

    def _page_count(self) -> int:
        if not self.filtered_rows:
            return 0
        return math.ceil(len(self.filtered_rows) / self._page_size())

    def _render_page(self) -> None:
        self.tree.delete(*self.tree.get_children())
        page_count = self._page_count()
        if page_count == 0:
            self.current_page = 0
            self.page_var.set("Page 0 of 0")
            self._update_pager_state()
            self.copy_button.state(["disabled"])
            return

        self.current_page = max(0, min(self.current_page, page_count - 1))
        page_size = self._page_size()
        start = self.current_page * page_size
        end = min(start + page_size, len(self.filtered_rows))

        for local_index, row in enumerate(self.filtered_rows[start:end]):
            global_row_number = start + local_index + 1
            values = [global_row_number] + [display_value(value) for value in row]
            tag = "even" if local_index % 2 == 0 else "odd"
            self.tree.insert("", "end", values=values, tags=(tag,))

        self.page_var.set(
            f"Page {self.current_page + 1:,} of {page_count:,} | Rows {start + 1:,}-{end:,} of {len(self.filtered_rows):,}"
        )
        self._update_pager_state()
        self.copy_button.state(["disabled"])

    def _page_size_changed(self) -> None:
        self.current_page = 0
        self._render_page()
        self._save_current_settings()

    def _first_page(self) -> None:
        self.current_page = 0
        self._render_page()

    def _previous_page(self) -> None:
        if self.current_page > 0:
            self.current_page -= 1
            self._render_page()

    def _next_page(self) -> None:
        if self.current_page + 1 < self._page_count():
            self.current_page += 1
            self._render_page()

    def _last_page(self) -> None:
        page_count = self._page_count()
        if page_count:
            self.current_page = page_count - 1
            self._render_page()

    def _update_pager_state(self) -> None:
        page_count = self._page_count()
        at_start = page_count == 0 or self.current_page == 0
        at_end = page_count == 0 or self.current_page >= page_count - 1
        for button in (self.first_button, self.previous_button):
            button.state(["disabled"] if at_start else ["!disabled"])
        for button in (self.next_button, self.last_button):
            button.state(["disabled"] if at_end else ["!disabled"])

    def _tree_selection_changed(self, _event: tk.Event[Any] | None = None) -> None:
        if self.tree.selection():
            self.copy_button.state(["!disabled"])
        else:
            self.copy_button.state(["disabled"])

    def _copy_selected(self) -> None:
        selected = self.tree.selection()
        if not selected:
            return
        lines = ["\t".join(("#", *self.columns))]
        for item_id in selected:
            values = self.tree.item(item_id, "values")
            lines.append("\t".join(str(value) for value in values))
        self.clipboard_clear()
        self.clipboard_append("\n".join(lines))
        self.status_var.set(f"Copied {len(selected):,} selected rows to the clipboard.")

    def _show_tree_context_menu(self, event: tk.Event[Any]) -> None:
        item_id = self.tree.identify_row(event.y)
        if item_id and item_id not in self.tree.selection():
            self.tree.selection_set(item_id)
        menu = tk.Menu(self, tearoff=False)
        menu.add_command(label="Copy selected rows", command=self._copy_selected)
        menu.add_command(label="Export current view", command=self._export_view)
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    def _export_view(self) -> None:
        if not self.filtered_rows:
            messagebox.showinfo("No Data", "There are no rows to export.", parent=self)
            return

        suggested_name = f"odw_well_monthly_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        output_name = filedialog.asksaveasfilename(
            parent=self,
            title="Export Current View",
            initialdir=self.last_export_dir,
            initialfile=suggested_name,
            defaultextension=".xlsx",
            filetypes=[
                ("Excel workbook", "*.xlsx"),
                ("CSV file", "*.csv"),
            ],
        )
        if not output_name:
            return

        output_path = Path(output_name)
        if output_path.suffix.lower() not in {".xlsx", ".csv"}:
            output_path = output_path.with_suffix(".xlsx")

        self.last_export_dir = str(output_path.parent)
        columns = tuple(self.columns)
        rows_snapshot = tuple(self.filtered_rows)
        self.status_var.set(f"Exporting {len(rows_snapshot):,} rows...")
        self._start_worker("export", export_rows, output_path, columns, rows_snapshot)

    def _accept_export_result(self, result: ExportResult) -> None:
        self.status_var.set(
            f"Exported {result.row_count:,} rows in {result.elapsed_seconds:,.1f} seconds."
        )
        self._save_current_settings()
        messagebox.showinfo(
            "Export Complete",
            f"Exported {result.row_count:,} rows to:\n\n{result.output_path}",
            parent=self,
        )

    def _show_sql(self) -> None:
        try:
            engineering_strings = parse_engineering_strings(self.engineering_var.get())
        except ValueError:
            engineering_strings = parse_engineering_strings(DEFAULT_ENGINEERING_STRINGS)
        sql, _params = build_query(
            engineering_strings,
            date(2020, 5, 1),
            date(2025, 4, 30),
        )

        window = tk.Toplevel(self)
        window.title("Parameterized SQL")
        window.geometry("980x720")
        window.minsize(700, 500)
        window.transient(self)

        frame = ttk.Frame(window, padding=10)
        frame.pack(fill="both", expand=True)
        frame.rowconfigure(1, weight=1)
        frame.columnconfigure(0, weight=1)

        ttk.Label(
            frame,
            text=(
                "Engineering strings and dates are passed as Oracle bind values. "
                "The UI's inclusive end date is converted to the next day for the '<' predicate."
            ),
            wraplength=920,
        ).grid(row=0, column=0, sticky="ew", pady=(0, 8))

        text_frame = ttk.Frame(frame)
        text_frame.grid(row=1, column=0, sticky="nsew")
        text_frame.rowconfigure(0, weight=1)
        text_frame.columnconfigure(0, weight=1)
        text_widget = tk.Text(
            text_frame,
            wrap="none",
            font=("Consolas", 10),
            background="#FBFCFE",
            foreground=self.COLORS["text"],
            padx=8,
            pady=8,
        )
        y_scroll = ttk.Scrollbar(text_frame, orient="vertical", command=text_widget.yview)
        x_scroll = ttk.Scrollbar(text_frame, orient="horizontal", command=text_widget.xview)
        text_widget.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)
        text_widget.grid(row=0, column=0, sticky="nsew")
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll.grid(row=1, column=0, sticky="ew")
        text_widget.insert("1.0", sql)
        text_widget.configure(state="disabled")

        ttk.Button(frame, text="Close", command=window.destroy).grid(
            row=2, column=0, sticky="e", pady=(8, 0)
        )

    def _settings_payload(self) -> dict[str, Any]:
        return {
            "username": self.username_var.get().strip(),
            "dsn": self.dsn_var.get().strip(),
            "engineering_strings": self.engineering_var.get().strip(),
            "start_date": self.start_date_var.get().strip(),
            "end_date": self.end_date_var.get().strip(),
            "max_rows": self.max_rows_var.get().strip(),
            "page_size": self.page_size_var.get().strip(),
            "last_export_dir": self.last_export_dir,
            "geometry": self.geometry(),
        }

    def _save_current_settings(self) -> None:
        # Password is intentionally excluded.
        save_settings(self._settings_payload())

    def _on_close(self) -> None:
        if self.busy_operation is not None:
            should_close = messagebox.askyesno(
                "Operation in Progress",
                f"A {self.busy_operation} operation is still running. Close the application anyway?",
                parent=self,
            )
            if not should_close:
                return
        self._save_current_settings()
        self.destroy()


def main() -> None:
    driver_mode = initialize_oracle_driver()
    app = ODWWellMonthlyViewer(driver_mode=driver_mode)
    app.mainloop()


if __name__ == "__main__":
    main()